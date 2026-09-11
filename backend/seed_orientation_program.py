"""One-off seed: parse "20 to CO.xlsx" (2nd Officer -> Chief Officer mentoring
checklist) into an OrientationProgram + its OrientationTask rows.

Idempotent-ish: if a program with this title already exists, its tasks are
replaced (matches the non-destructive-upsert spirit of app/seed.py, but
simpler since orientation tasks have no learner state attached directly —
OrientationTaskCompletion rows key off task_id, so re-running after
enrollments exist would orphan completions; only safe to re-run before any
enrollment is created).

Run from backend/:  ./venv/Scripts/python.exe seed_orientation_program.py [path-to-xlsx]
"""
import re
import sys
import os

sys.path.insert(0, os.path.dirname(__file__))

import openpyxl
from app.database import SessionLocal
from app import models

DEFAULT_XLSX = r"C:\Users\sajin\Downloads\20 to CO.xlsx"

PROGRAM_TITLE = "2nd Officer to Chief Officer"
PROGRAM_SUBTITLE = "Deck mentoring checklist"
DEPARTMENT = "deck"
FROM_RANK = "Second Officer"
TO_RANK = "Chief Officer"


def parse_tasks(xlsx_path: str):
    wb = openpyxl.load_workbook(xlsx_path)
    ws = wb.active
    values = [row[0] for row in ws.iter_rows(min_row=3, max_row=ws.max_row, values_only=True) if row[0]]
    full_text = "\n".join(str(v) for v in values)
    blocks = re.split(r"(?=Task Name)", full_text)

    tasks = []
    for block in blocks:
        if "Task Name" not in block:
            continue
        m = re.search(r"Task Name\s+(.*?)\s*\n\s*Task Description\s+(.*)", block, re.S)
        if not m:
            continue
        title = m.group(1).strip()
        desc = m.group(2)
        desc = re.sub(r"\(Relevant task supporting documents.*", "", desc, flags=re.S).strip()
        desc = re.sub(r"MASTER'?S? Initials:.*", "", desc, flags=re.S | re.I).strip()
        desc = re.sub(r"Ref\. Code.*", "", desc, flags=re.S).strip()
        desc = re.sub(r"Page \d+ of \d+", "", desc).strip()
        if title:
            tasks.append((title, desc))
    return tasks


def run(xlsx_path: str = DEFAULT_XLSX):
    tasks = parse_tasks(xlsx_path)
    print(f"Parsed {len(tasks)} tasks from {xlsx_path}")

    db = SessionLocal()
    try:
        program = db.query(models.OrientationProgram).filter_by(title=PROGRAM_TITLE).first()
        if program:
            if program.enrollments:
                print(f"'{PROGRAM_TITLE}' already has {len(program.enrollments)} enrollment(s) — "
                      "refusing to replace its tasks (would orphan completions). Aborting.")
                return
            print(f"'{PROGRAM_TITLE}' already exists with no enrollments — replacing its tasks.")
            db.query(models.OrientationTask).filter_by(program_id=program.id).delete()
        else:
            program = models.OrientationProgram(
                title=PROGRAM_TITLE, subtitle=PROGRAM_SUBTITLE, department=DEPARTMENT,
                from_rank=FROM_RANK, to_rank=TO_RANK, is_active=True, order=0,
            )
            db.add(program)
            db.flush()

        for i, (title, desc) in enumerate(tasks):
            db.add(models.OrientationTask(program_id=program.id, title=title, description=desc, order=i))

        db.commit()
        print(f"Seeded program '{PROGRAM_TITLE}' (id={program.id}) with {len(tasks)} tasks.")
    finally:
        db.close()


if __name__ == "__main__":
    path = sys.argv[1] if len(sys.argv) > 1 else DEFAULT_XLSX
    run(path)
