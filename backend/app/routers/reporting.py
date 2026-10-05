"""Compliance reporting, analytics dashboard KPIs, and Excel/CSV export router."""
import csv
import io
from collections import defaultdict
from datetime import datetime, timezone, timedelta, date
from typing import Optional

from fastapi import APIRouter, Depends, HTTPException, Response
from sqlalchemy.orm import Session
from sqlalchemy import func

from app.database import get_db
from app import models
from app.auth import require_admin, get_current_user
from app.course_helpers import (
    _maybe_auto_complete_course, enrolled_course_ids, get_progress
)

router = APIRouter(tags=["Reporting"])


def _report(db):
    courses = db.query(models.Course).order_by(models.Course.order).all()
    learners = (db.query(models.User).filter_by(role="learner")
                .order_by(models.User.full_name).all())
    
    enrollments = db.query(models.Enrollment).all()
    progress_all = db.query(models.Progress).all()
    certs = db.query(models.Certificate).all()
    approvals = (db.query(models.AssessmentApproval)
                 .order_by(models.AssessmentApproval.created_at).all())
    attempts_counts = db.query(models.Attempt.learner_id, models.Attempt.course_id, func.count(models.Attempt.id)).group_by(models.Attempt.learner_id, models.Attempt.course_id).all()

    e_map = {(e.learner_id, e.course_id): e for e in enrollments}
    p_map = {(p.learner_id, p.course_id): p for p in progress_all}
    c_map = {(c.learner_id, c.course_id): c for c in certs}
    ap_map = {}
    for a in approvals:
        ap_map[(a.learner_id, a.course_id)] = a
    a_map = {(l, c): count for l, c, count in attempts_counts}
    
    course_chapter_ids = {c.id: {ch.id for ch in c.chapters} for c in courses}
    course_chapter_counts = {cid: len(ids) for cid, ids in course_chapter_ids.items()}
    
    rows = []
    for lr in learners:
        assigned_courses = [c for c in courses if (lr.id, c.id) in e_map]
        cells = {}
        for c in assigned_courses:
            enr = e_map.get((lr.id, c.id))
            prog = p_map.get((lr.id, c.id))
            cert = c_map.get((lr.id, c.id))
            ap = ap_map.get((lr.id, c.id))
            attempts = a_map.get((lr.id, c.id), 0)
            
            if prog and prog.passed:
                status = "passed"
            elif prog and ((prog.completed_chapters and len(prog.completed_chapters) > 0) or prog.score is not None):
                status = "in-progress"
            else:
                status = "assigned"
            
            total_chs = course_chapter_counts.get(c.id, 0)
            valid_ids = course_chapter_ids.get(c.id, set())
            done_chs  = len(valid_ids & set(prog.completed_chapters or [])) if prog else 0
            pct       = round(done_chs / total_chs * 100) if total_chs else 0
            
            if cert and cert.issued_at:
                passed_on = cert.issued_at.strftime("%Y-%m-%d")
            elif ap and ap.created_at:
                passed_on = ap.created_at.strftime("%Y-%m-%d")
            else:
                passed_on = None
            
            cells[c.id] = {
                "status": status,
                "score": prog.score if prog else None,
                "startedOn": enr.assigned_at.strftime("%Y-%m-%d") if enr and enr.assigned_at else None,
                "passedOn": passed_on,
                "pendingApproval": bool(ap and ap.status == "pending"),
                "approvalRemark": ap.remark if ap and ap.remark else None,
                "attempts": attempts,
                "completionPct":     pct,
                "completedChapters": done_chs,
                "totalChapters":     total_chs,
            }
        
        rows.append({
            "learnerId": lr.id, "name": lr.full_name, "crewId": lr.crew_id,
            "rank": lr.rank, "mobileNo": lr.mobile_number,
            "isActive": bool(lr.is_active), "cells": cells,
        })
    return courses, rows


@router.get("/api/admin/report")
def admin_report(admin: models.User = Depends(require_admin), db: Session = Depends(get_db)):
    courses, rows = _report(db)
    return {"courses": [{"id": c.id, "title": c.title} for c in courses], "rows": rows}


@router.post("/api/admin/maintenance/reconcile-progress")
def admin_reconcile_progress(admin: models.User = Depends(require_admin), db: Session = Depends(get_db)):
    """One-off heal for learners stuck from past chapter deletions."""
    healed = []
    for course in db.query(models.Course).all():
        for p in db.query(models.Progress).filter_by(course_id=course.id).all():
            if _maybe_auto_complete_course(db, course, p, p.learner_id):
                learner = db.get(models.User, p.learner_id)
                healed.append({
                    "courseId": course.id, "courseTitle": course.title,
                    "learnerId": p.learner_id,
                    "learnerName": learner.full_name if learner else None,
                })
    db.commit()
    return {"healedCount": len(healed), "healed": healed}


@router.get("/api/admin/dashboard-stats")
def admin_dashboard_stats(admin: models.User = Depends(require_admin), db: Session = Depends(get_db)):
    """Rich analytics for the admin dashboard charts."""
    learners = db.query(models.User).filter_by(role="learner").all()
    rank_counts = defaultdict(int)
    status_counts = defaultdict(int)
    vessel_counts = defaultdict(int)
    for lr in learners:
        rank_counts[lr.rank or "Unspecified"] += 1
        status_counts[lr.emp_status or "Unknown"] += 1
        vessel_counts[lr.current_vessel or "Unassigned"] += 1

    rank_data = sorted(
        [{"rank": k, "count": v} for k, v in rank_counts.items()],
        key=lambda x: -x["count"]
    )[:12]

    status_data = [{"status": k, "count": v} for k, v in status_counts.items() if v > 0]

    vessel_data = sorted(
        [{"vessel": k, "count": v} for k, v in vessel_counts.items() if k != "Unassigned"],
        key=lambda x: -x["count"]
    )[:8]

    courses = db.query(models.Course).order_by(models.Course.order).all()
    enrollments = db.query(models.Enrollment).all()
    progress_all = db.query(models.Progress).all()
    
    e_map = defaultdict(set)
    for e in enrollments:
        e_map[e.course_id].add(e.learner_id)
        
    p_map = {(p.learner_id, p.course_id): p for p in progress_all}
    
    course_stats = []
    for c in courses:
        enrolled_ids = e_map[c.id]
        total = len(enrolled_ids)
        if total == 0:
            course_stats.append({
                "course": c.title[:30], "courseId": c.id,
                "enrolled": 0, "passed": 0, "inProgress": 0,
                "assigned": 0, "passRate": 0,
            })
            continue
        passed = inprogress = 0
        for lid in enrolled_ids:
            prog = p_map.get((lid, c.id))
            if prog and prog.passed:
                passed += 1
            elif prog and ((prog.completed_chapters and len(prog.completed_chapters) > 0) or prog.score is not None):
                inprogress += 1
        assigned = total - passed - inprogress
        course_stats.append({
            "course": c.title[:30], "courseId": c.id,
            "enrolled": total, "passed": passed,
            "inProgress": inprogress, "assigned": assigned,
            "passRate": round((passed / total) * 100) if total else 0,
        })

    today = date.today()
    months = []
    for i in range(5, -1, -1):
        month_date = today.replace(day=1)
        for _ in range(i):
            month_date = (month_date - timedelta(days=1)).replace(day=1)
        months.append(month_date)

    enrollment_trend = []
    for m in months:
        if m.month == 12:
            next_m = m.replace(year=m.year + 1, month=1)
        else:
            next_m = m.replace(month=m.month + 1)
        count = (db.query(models.Enrollment)
                 .filter(models.Enrollment.assigned_at >= datetime(m.year, m.month, 1, tzinfo=timezone.utc),
                         models.Enrollment.assigned_at < datetime(next_m.year, next_m.month, 1, tzinfo=timezone.utc))
                 .count())
        enrollment_trend.append({"month": m.strftime("%b %Y"), "enrollments": count})

    recent_certs = (db.query(models.Certificate)
                    .order_by(models.Certificate.issued_at.desc())
                    .limit(8).all())
    recent_cert_list = []
    for cert in recent_certs:
        user_obj = db.get(models.User, cert.learner_id)
        course_obj = db.query(models.Course).filter_by(id=cert.course_id).first()
        recent_cert_list.append({
            "id": cert.id,
            "learner": user_obj.full_name if user_obj else "Unknown",
            "rank": user_obj.rank if user_obj else None,
            "course": course_obj.title if course_obj else cert.course_id,
            "score": cert.score,
            "issuedAt": cert.issued_at.isoformat() if cert.issued_at else None,
        })

    total_crew = len(learners)
    active_crew = sum(1 for u in learners if u.is_active)
    total_enrollments = db.query(models.Enrollment).count()
    total_certs = db.query(models.Certificate).count()
    total_courses = db.query(models.Course).count()
    total_attempts = db.query(models.Attempt).count()
    pass_attempts = db.query(models.Attempt).filter_by(passed=True).count()

    last_sync_log = (db.query(models.SyncLog)
                      .order_by(models.SyncLog.started_at.desc())
                      .first())
    last_sync = None
    if last_sync_log:
        last_sync = {
            "startedAt": last_sync_log.started_at.isoformat() if last_sync_log.started_at else None,
            "finishedAt": last_sync_log.finished_at.isoformat() if last_sync_log.finished_at else None,
            "status": last_sync_log.status,
            "recordsFetched": last_sync_log.records_fetched,
            "recordsCreated": last_sync_log.records_created,
            "recordsUpdated": last_sync_log.records_updated,
            "errorMessage": last_sync_log.error_message,
        }

    return {
        "kpis": {
            "totalCrew": total_crew,
            "activeCrew": active_crew,
            "totalEnrollments": total_enrollments,
            "totalCertificates": total_certs,
            "totalCourses": total_courses,
            "totalAttempts": total_attempts,
            "passAttempts": pass_attempts,
            "overallPassRate": round((pass_attempts / total_attempts) * 100) if total_attempts else 0,
        },
        "crewByRank": rank_data,
        "crewByStatus": status_data,
        "crewByVessel": vessel_data,
        "courseStats": course_stats,
        "enrollmentTrend": enrollment_trend,
        "recentCertificates": recent_cert_list,
        "lastSync": last_sync,
    }


@router.get("/api/admin/report.csv")
def admin_report_csv(admin: models.User = Depends(require_admin), db: Session = Depends(get_db)):
    courses, rows = _report(db)
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow(["Crew ID", "Name", "Rank", "Active", "Course", "Status", "Score", "Completed On", "Approval Remark"])
    for r in rows:
        active = "yes" if r["isActive"] else "no"
        if not r["cells"]:
            w.writerow([r["crewId"], r["name"], r["rank"] or "", active,
                        "(no courses assigned)", "", "", "", ""])
            continue
        for c in courses:
            cell = r["cells"].get(c.id)
            if not cell:
                continue
            w.writerow([r["crewId"], r["name"], r["rank"] or "", active, c.title,
                        cell["status"], "" if cell["score"] is None else cell["score"],
                        cell["passedOn"] or "", cell.get("approvalRemark") or ""])
    return Response(content=buf.getvalue(), media_type="text/csv",
                    headers={"Content-Disposition": "attachment; filename=ozellar-compliance-report.csv"})


@router.get("/api/admin/report.xlsx")
def admin_report_xlsx(
    admin: models.User = Depends(require_admin),
    db: Session = Depends(get_db),
    crew_search: str = None,
    course_id:   str = None,
    status:      str = None,
):
    """Download the compliance report as a styled Excel workbook."""
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter

    courses, rows = _report(db)

    if crew_search:
        q = crew_search.strip().lower()
        rows = [r for r in rows if
                q in (r["name"] or "").lower() or
                q in (r["crewId"] or "").lower() or
                q in (r["rank"] or "").lower()]

    if course_id:
        rows = [r for r in rows if course_id in r["cells"]]
        courses = [c for c in courses if c.id == course_id]

    if status:
        if course_id:
            rows = [r for r in rows if r["cells"].get(course_id, {}).get("status") == status]
        else:
            rows = [r for r in rows
                    if any(cell["status"] == status for cell in r["cells"].values())]

    filter_note = []
    if crew_search: filter_note.append(f'Crew: "{crew_search}"')
    if course_id:
        c_obj = next((c for c in courses if c.id == course_id), None)
        if c_obj: filter_note.append(f"Course: {c_obj.title}")
    if status:
        filter_note.append(f"Status: {status}")

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Compliance Report"

    HDR_FILL   = PatternFill("solid", fgColor="1E3A5F")
    PASS_FILL  = PatternFill("solid", fgColor="D6F5E3")
    WIP_FILL   = PatternFill("solid", fgColor="FFF3CD")
    ASGN_FILL  = PatternFill("solid", fgColor="F2F4F7")
    WHITE_FILL = PatternFill("solid", fgColor="FFFFFF")
    HDR_FONT   = Font(bold=True, color="FFFFFF", size=11)
    BODY_FONT  = Font(size=10)
    BOLD_FONT  = Font(bold=True, size=10)
    thin = Side(style="thin", color="D0D5DD")
    thick_bottom = Side(style="medium", color="475467")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    border_thick_bottom = Border(left=thin, right=thin, top=thin, bottom=thick_bottom)
    center = Alignment(horizontal="center", vertical="center")
    left   = Alignment(horizontal="left",   vertical="center", wrap_text=False)

    title_text = "Ozellar Marine — Compliance Report"
    if filter_note:
        title_text += f"  |  Filters: {', '.join(filter_note)}"
    ws.merge_cells("A1:K1")
    tc = ws.cell(row=1, column=1, value=title_text)
    tc.font = Font(bold=True, size=12, color="1E3A5F")
    tc.alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[1].height = 22

    ws.merge_cells("A2:K2")
    sc = ws.cell(row=2, column=1,
                 value=f"Generated: {datetime.now(timezone.utc).strftime('%d %b %Y %H:%M UTC')}  ·  "
                       f"{len(rows)} crew member(s) shown")
    sc.font = Font(size=9, color="5C626D")
    sc.alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[2].height = 16

    headers = ["Crew ID", "Name", "Rank", "Active",
               "Course", "Status", "Score (%)", "Attempts", "Started On", "Completed On", "Approval Remark"]
    ws.append(headers)
    for ci, h in enumerate(headers, 1):
        cell = ws.cell(row=3, column=ci)
        cell.fill = HDR_FILL
        cell.font = HDR_FONT
        cell.alignment = center
        cell.border = border
    ws.row_dimensions[3].height = 22

    status_labels = {"passed": "Completed", "in-progress": "In Progress", "assigned": "Not Started"}
    row_idx = 4
    for r in rows:
        active = "Yes" if r["isActive"] else "No"
        start_row = row_idx
        
        if not r["cells"]:
            data = [r["crewId"], r["name"], r["rank"] or "", active,
                    "(no courses assigned)", "", "", "", "", "", ""]
            ws.append(data)
            for ci in range(1, 12):
                c = ws.cell(row=row_idx, column=ci)
                c.fill = ASGN_FILL if ci > 4 else WHITE_FILL
                c.font = BODY_FONT
                c.alignment = left; c.border = border
            row_idx += 1
            continue

        for course in courses:
            cell_data = r["cells"].get(course.id)
            if not cell_data:
                continue
            status = cell_data["status"]
            label  = status_labels.get(status, status)
            score  = cell_data["score"] if cell_data["score"] is not None else ""
            started_on = cell_data.get("startedOn") or ""
            passed_on = cell_data.get("passedOn") or ""
            attempts = cell_data.get("attempts", 0)
            if attempts == 0: attempts = ""
            remark = cell_data.get("approvalRemark") or ""

            fill = PASS_FILL if status == "passed" else (WIP_FILL if status == "in-progress" else ASGN_FILL)
            data = [r["crewId"], r["name"], r["rank"] or "", active,
                    course.title, label, score, attempts, started_on, passed_on, remark]
            ws.append(data)
            for ci, val in enumerate(data, 1):
                c = ws.cell(row=row_idx, column=ci)
                c.fill = fill if ci > 4 else WHITE_FILL
                c.font = BOLD_FONT if ci == 2 else BODY_FONT
                c.alignment = left
                c.border = border
            row_idx += 1

        if row_idx - 1 > start_row:
            for ci in range(1, 5):
                ws.merge_cells(start_row=start_row, start_column=ci, end_row=row_idx-1, end_column=ci)

        for ci in range(1, 12):
            ws.cell(row=row_idx-1, column=ci).border = border_thick_bottom

    col_widths = [14, 26, 18, 8, 36, 14, 11, 10, 14, 14, 40]
    for ci, w in enumerate(col_widths, 1):
        ws.column_dimensions[get_column_letter(ci)].width = w

    ws.freeze_panes = "A4"

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return Response(
        content=buf.read(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=ozellar-compliance-report.xlsx"},
    )


@router.get("/api/crew/my-report.xlsx")
def crew_my_report_xlsx(status: Optional[str] = None, user: models.User = Depends(get_current_user), db: Session = Depends(get_db)):
    """Crew personal training record as a styled Excel workbook."""
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter

    if user.role in ("admin", "super_admin"):
        raise HTTPException(403, "Use /api/admin/report.xlsx for admin reports")

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "My Training Record"

    HDR_FILL  = PatternFill("solid", fgColor="1E3A5F")
    PASS_FILL = PatternFill("solid", fgColor="D6F5E3")
    WIP_FILL  = PatternFill("solid", fgColor="FFF3CD")
    ASGN_FILL = PatternFill("solid", fgColor="F2F4F7")
    HDR_FONT  = Font(bold=True, color="FFFFFF", size=11)
    BODY_FONT = Font(size=10)
    BOLD_FONT  = Font(bold=True, size=10)
    thin   = Side(style="thin", color="D0D5DD")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    center = Alignment(horizontal="center", vertical="center")
    left   = Alignment(horizontal="left",   vertical="center")

    ws.merge_cells("A1:L1")
    title_cell = ws.cell(row=1, column=1,
                         value=f"Training Record — {user.full_name}  ·  {user.rank or ''}")
    title_cell.font = Font(bold=True, size=13, color="1E3A5F")
    title_cell.alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[1].height = 26

    ws.merge_cells("A2:L2")
    sub_cell = ws.cell(row=2, column=1,
                       value=f"Crew ID: {user.crew_id or 'N/A'}   |   Generated: {datetime.now(timezone.utc).strftime('%d %b %Y')}")
    sub_cell.font = Font(size=10, color="5C626D")
    ws.row_dimensions[2].height = 18

    headers = ["Course", "Status", "Attempts", "Chapters Done", "Total Chapters",
               "Progress (%)", "Score (%)", "Grade", "Started On", "Completed On",
               "Time Taken (days)", "Certificate ID"]
    ws.append(headers)
    for ci, h in enumerate(headers, 1):
        c = ws.cell(row=3, column=ci)
        c.fill = HDR_FILL; c.font = HDR_FONT
        c.alignment = center; c.border = border
    ws.row_dimensions[3].height = 22

    assigned_ids = enrolled_course_ids(db, user.id)
    courses = (db.query(models.Course).filter(models.Course.id.in_(assigned_ids))
               .order_by(models.Course.order).all()) if assigned_ids else []

    row_idx = 4
    total_passed = 0
    for c in courses:
        prog = get_progress(db, user.id, c.id)
        cert = (db.query(models.Certificate)
                .filter_by(learner_id=user.id, course_id=c.id).first())

        valid_ids   = {ch.id for ch in c.chapters}
        done_count  = len(valid_ids & set(prog.completed_chapters or [])) if prog else 0
        total_ch    = len(c.chapters)
        pct         = round(done_count / total_ch * 100) if total_ch else 0
        score       = prog.score if prog else None
        passed_flag = bool(prog.passed) if prog else False

        if passed_flag:
            if not cert:
                calc_status = "Pending Approval"
                fill = WIP_FILL
            else:
                calc_status = "Completed"
                fill = PASS_FILL
                total_passed += 1
        elif prog and (done_count > 0 or score is not None):
            calc_status = "In Progress"
            fill = WIP_FILL
        else:
            calc_status = "Not Started"
            fill = ASGN_FILL

        if status == 'completed' and calc_status != "Completed": continue
        if status == 'in-progress' and calc_status not in ["In Progress", "Pending Approval"]: continue
        if status == 'not-started' and calc_status != "Not Started": continue

        grade = ""
        if score is not None:
            if score >= 90:   grade = "A+"
            elif score >= 80: grade = "A"
            elif score >= 70: grade = "B"
            elif score >= 60: grade = "C"
            else:             grade = "F"

        time_days = ""
        started_on = ""
        enr = (db.query(models.Enrollment)
               .filter_by(learner_id=user.id, course_id=c.id).first())
        if enr and enr.assigned_at:
            started_on = enr.assigned_at.strftime("%d %b %Y")
            if cert and cert.issued_at:
                delta = cert.issued_at - enr.assigned_at
                time_days = max(0, delta.days)

        attempts = (db.query(models.Attempt)
                    .filter_by(learner_id=user.id, course_id=c.id).count())
        if attempts == 0: attempts = ""

        cert_id   = cert.id if cert else ""
        if passed_flag and not cert:
            passed_on = "Pending Approval"
        else:
            passed_on = cert.issued_at.strftime("%d %b %Y") if cert and cert.issued_at else ""
        score_str = score if score is not None else ""

        row_data = [c.title, calc_status, attempts, done_count, total_ch,
                    pct, score_str, grade, started_on, passed_on, time_days, cert_id]
        ws.append(row_data)

        for ci, val in enumerate(row_data, 1):
            cell = ws.cell(row=row_idx, column=ci)
            cell.fill = fill
            cell.font = BOLD_FONT if ci == 1 else BODY_FONT
            cell.alignment = left
            cell.border = border
        row_idx += 1

    ws.append([])
    row_idx += 1
    summary_row = row_idx
    ws.cell(row=summary_row, column=1, value="SUMMARY").font = Font(bold=True, size=10, color="1E3A5F")
    ws.cell(row=summary_row, column=2,
            value=f"{total_passed} of {len(courses)} courses passed").font = BODY_FONT

    col_widths = [36, 15, 10, 15, 15, 14, 11, 8, 14, 14, 18, 22]
    for ci, w in enumerate(col_widths, 1):
        ws.column_dimensions[get_column_letter(ci)].width = w

    ws.freeze_panes = "A4"

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return Response(
        content=buf.read(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename=ozellar-training-{user.crew_id or user.id}.xlsx"},
    )
