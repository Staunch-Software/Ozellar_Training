"""Screening entrance test router: candidate exam session + admin test builder & reporting."""
import io
import os
import re
import uuid
from datetime import datetime, timezone
from typing import Optional

from fastapi import APIRouter, Depends, HTTPException, UploadFile, File, Response
from fastapi.responses import StreamingResponse
from pydantic import BaseModel
from sqlalchemy.orm import Session
from sqlalchemy import func

from app.config import UPLOAD_DIR
from app.database import get_db
from app import models, storage
from app.auth import (
    create_screening_token, get_current_candidate, candidate_public,
    require_admin, verify_password, hash_password, normalize_name,
    check_rate_limit, clear_rate_limit
)
from app.security_upload import validate_safe_upload

router = APIRouter(tags=["Screening"])

SCREENING_PHOTO_DIR = os.path.join(UPLOAD_DIR, "screening_photos")
os.makedirs(SCREENING_PHOTO_DIR, exist_ok=True)


class ScreeningLoginRequest(BaseModel):
    name: str
    password: str


class ScreeningSubmitRequest(BaseModel):
    personal_data: Optional[dict] = None
    section_answers: dict  # {section_id: [int | null]}


class CreateTestRequest(BaseModel):
    title: str
    timerMinutes: int = 80
    correctScore: int = 4
    wrongPenalty: int = 1


class SectionRequest(BaseModel):
    id: Optional[str] = None
    title: str
    section_type: str = "mcq"
    passage: Optional[str] = None
    order: int = 0


class QuestionRequest(BaseModel):
    id: Optional[str] = None
    prompt: str
    options: list[str]
    answer: int
    imageUrls: Optional[list[str]] = None
    order: int = 0


class BulkQuestionsRequest(BaseModel):
    questions: list[QuestionRequest]


class CreateCandidateRequest(BaseModel):
    fullName: str
    password: str
    testId: str
    mobileNumber: Optional[str] = None


def serialize_test(t, include_sections=False, db=None):
    candidate_count = len(t.candidates) if t.candidates else 0
    submitted_count = sum(1 for c in (t.candidates or []) if c.status == "submitted")
    total_questions = sum(
        len(s.questions) for s in (t.sections or []) if s.section_type == "mcq"
    )
    data = {
        "id": t.id, "title": t.title, "timerMinutes": t.timer_minutes,
        "correctScore": t.correct_score, "wrongPenalty": t.wrong_penalty,
        "isActive": t.is_active, "createdAt": t.created_at.isoformat() if t.created_at else None,
        "candidateCount": candidate_count, "submittedCount": submitted_count,
        "totalQuestions": total_questions,
    }
    if include_sections:
        data["sections"] = [
            {
                "id": s.id, "title": s.title, "type": s.section_type,
                "passage": s.passage, "order": s.order,
                "questions": [
                    {"id": q.id, "prompt": q.prompt, "options": q.options,
                     "answer": q.answer, "imageUrls": q.image_urls or [], "order": q.order}
                    for q in s.questions
                ],
            }
            for s in (t.sections or [])
        ]
    return data


# --- Candidate Endpoints ---

@router.post("/api/screening/login")
def screening_login(req: ScreeningLoginRequest, db: Session = Depends(get_db)):
    """Test-taker login: full name + password."""
    name = (req.name or "").strip()
    if not name or not req.password:
        raise HTTPException(401, "Name and password are required")
    check_rate_limit(db, f"screening:{normalize_name(name)}")
    candidates = db.query(models.ScreeningCandidate).filter_by(is_active=True).all()
    target = normalize_name(name)
    match = None
    for c in candidates:
        if normalize_name(c.full_name) == target:
            if verify_password(req.password, c.password_hash):
                match = c
                break
    if not match:
        raise HTTPException(401, "Invalid name or password")
    if match.status == "submitted":
        raise HTTPException(403, "You have already submitted this test")
    clear_rate_limit(db, f"screening:{normalize_name(name)}")
    token = create_screening_token(match)
    return {"token": token, "candidate": candidate_public(match)}


@router.get("/api/screening/test")
def screening_get_test(
    candidate: models.ScreeningCandidate = Depends(get_current_candidate),
    db: Session = Depends(get_db),
):
    """Get the full test details for the current candidate (no answer keys)."""
    test = db.get(models.ScreeningTest, candidate.test_id)
    if not test or not test.is_active:
        raise HTTPException(404, "Test not found or inactive")
    attempt = candidate.attempt
    sections_data = []
    for sec in test.sections:
        s = {
            "id": sec.id, "title": sec.title,
            "type": sec.section_type, "order": sec.order,
            "passage": sec.passage,
            "questions": [
                {"id": q.id, "prompt": q.prompt, "options": q.options,
                 "imageUrls": q.image_urls or [], "order": q.order}
                for q in sec.questions
            ],
        }
        sections_data.append(s)

    remaining_seconds = None
    if attempt and attempt.started_at and not attempt.submitted_at:
        db_now = db.query(func.now()).scalar()
        started = attempt.started_at
        if db_now.tzinfo is not None:
            db_now = db_now.replace(tzinfo=None)
        if started.tzinfo is not None:
            started = started.replace(tzinfo=None)
        elapsed = (db_now - started).total_seconds()
        remaining_seconds = max(0, int(test.timer_minutes * 60 - elapsed))

    return {
        "id": test.id, "title": test.title,
        "timerMinutes": test.timer_minutes,
        "correctScore": test.correct_score,
        "wrongPenalty": test.wrong_penalty,
        "remainingSeconds": remaining_seconds,
        "sections": sections_data,
        "attempt": {
            "startedAt": attempt.started_at.isoformat() if attempt else None,
            "submittedAt": attempt.submitted_at.isoformat() if attempt and attempt.submitted_at else None,
            "status": candidate.status,
            "tabSwitchCount": attempt.tab_switch_count or 0,
            "sectionAnswers": attempt.section_answers or {},
            "personalData": attempt.personal_data or {},
        } if attempt else None,
    }


@router.post("/api/screening/start")
def screening_start(
    candidate: models.ScreeningCandidate = Depends(get_current_candidate),
    db: Session = Depends(get_db),
):
    """Create the attempt row (records start time). Idempotent."""
    if candidate.status == "submitted":
        raise HTTPException(403, "Test already submitted")
    existing = candidate.attempt
    if existing:
        return {"startedAt": existing.started_at.isoformat(), "ok": True}
    attempt = models.ScreeningAttempt(
        candidate_id=candidate.id, test_id=candidate.test_id,
    )
    db.add(attempt)
    candidate.status = "in_progress"
    db.commit()
    db.refresh(attempt)
    return {"startedAt": attempt.started_at.isoformat(), "ok": True}


@router.post("/api/screening/tab-switch")
def screening_tab_switch(
    candidate: models.ScreeningCandidate = Depends(get_current_candidate),
    db: Session = Depends(get_db),
):
    """Record one instance of the candidate leaving the exam tab/window."""
    if candidate.status == "submitted":
        raise HTTPException(403, "Test already submitted")
    attempt = candidate.attempt
    if not attempt:
        raise HTTPException(400, "Test not started")
    attempt.tab_switch_count = (attempt.tab_switch_count or 0) + 1
    db.commit()
    return {"tabSwitchCount": attempt.tab_switch_count}


@router.post("/api/screening/photo")
async def screening_upload_photo(
    file: UploadFile = File(...),
    candidate: models.ScreeningCandidate = Depends(get_current_candidate),
    db: Session = Depends(get_db),
):
    """Upload passport-size photo for the candidate."""
    from PIL import Image
    await validate_safe_upload(file)
    data = await file.read()
    try:
        img = Image.open(io.BytesIO(data))
        img.verify()
    except Exception:
        raise HTTPException(400, "Invalid image file")
    dest = os.path.join(SCREENING_PHOTO_DIR, f"{candidate.id}.jpg")
    img2 = Image.open(io.BytesIO(data)).convert("RGB")
    img2.thumbnail((600, 600))
    img2.save(dest, "JPEG", quality=85)
    if candidate.attempt:
        candidate.attempt.photo_path = dest
    db.commit()
    return {"ok": True, "hasPhoto": True}


@router.get("/api/screening/photo")
def screening_get_photo(
    candidate: models.ScreeningCandidate = Depends(get_current_candidate),
):
    """Serve the signed-in candidate's own identity photo."""
    path = os.path.join(SCREENING_PHOTO_DIR, f"{candidate.id}.jpg")
    if not os.path.exists(path):
        raise HTTPException(404, "No photo on record")
    with open(path, "rb") as fh:
        return Response(fh.read(), media_type="image/jpeg",
                        headers={"Cache-Control": "private, max-age=300"})


@router.post("/api/screening/submit")
def screening_submit(
    req: ScreeningSubmitRequest,
    candidate: models.ScreeningCandidate = Depends(get_current_candidate),
    db: Session = Depends(get_db),
):
    """Grade and finalise the test. Returns the scored result."""
    if candidate.status == "submitted":
        raise HTTPException(403, "Test already submitted")
    test = db.get(models.ScreeningTest, candidate.test_id)
    if not test:
        raise HTTPException(404, "Test not found")

    attempt = candidate.attempt
    if not attempt:
        attempt = models.ScreeningAttempt(candidate_id=candidate.id, test_id=candidate.test_id)
        db.add(attempt)
        db.flush()

    correct = wrong = unanswered = 0
    section_results = []
    for sec in test.sections:
        if sec.section_type == "personal_data":
            continue
        answers = req.section_answers.get(sec.id, [])
        sq = 0; swrong = 0; sunanswered = 0
        for i, q in enumerate(sec.questions):
            chosen = answers[i] if i < len(answers) else None
            if chosen is None:
                sunanswered += 1
            elif chosen == q.answer:
                sq += 1
            else:
                swrong += 1
        correct += sq
        wrong += swrong
        unanswered += sunanswered
        section_results.append({
            "sectionId": sec.id, "title": sec.title,
            "correct": sq, "wrong": swrong, "unanswered": sunanswered,
            "total": len(sec.questions),
        })

    raw_score = correct * test.correct_score - wrong * test.wrong_penalty

    attempt.personal_data = req.personal_data
    attempt.section_answers = req.section_answers
    attempt.score = raw_score
    attempt.correct_count = correct
    attempt.wrong_count = wrong
    attempt.unanswered_count = unanswered
    attempt.submitted_at = datetime.now(timezone.utc)

    candidate.status = "submitted"
    if req.personal_data:
        mob = req.personal_data.get("mobile") or req.personal_data.get("mobileNumber")
        if mob and not candidate.mobile_number:
            candidate.mobile_number = mob
    db.commit()
    return {
        "score": raw_score, "correct": correct, "wrong": wrong,
        "unanswered": unanswered,
        "correctScore": test.correct_score,
        "wrongPenalty": test.wrong_penalty,
        "sectionResults": section_results,
        "submittedAt": attempt.submitted_at.isoformat(),
    }


@router.post("/api/screening/autosave")
def screening_autosave(
    req: ScreeningSubmitRequest,
    candidate: models.ScreeningCandidate = Depends(get_current_candidate),
    db: Session = Depends(get_db),
):
    """Periodically persist in-progress answers."""
    if candidate.status == "submitted":
        raise HTTPException(403, "Test already submitted")
    attempt = candidate.attempt
    if not attempt:
        raise HTTPException(400, "Test not started")
    attempt.section_answers = req.section_answers
    if req.personal_data:
        attempt.personal_data = req.personal_data
    db.commit()
    return {"ok": True, "savedAt": datetime.now(timezone.utc).isoformat()}


@router.get("/api/screening/result")
def screening_result(
    candidate: models.ScreeningCandidate = Depends(get_current_candidate),
    db: Session = Depends(get_db),
):
    """Get the result for the current candidate (after submission)."""
    if candidate.status != "submitted":
        raise HTTPException(400, "Test not yet submitted")
    attempt = candidate.attempt
    if not attempt:
        raise HTTPException(404, "Attempt not found")
    test = db.get(models.ScreeningTest, candidate.test_id)

    total_questions = sum(
        len(s.questions) for s in (test.sections or [])
        if s.section_type != "personal_data"
    ) if test else 0
    max_score = total_questions * (test.correct_score if test else 4)

    time_taken_minutes = None
    if attempt.started_at and attempt.submitted_at:
        started, ended = attempt.started_at, attempt.submitted_at
        if (started.tzinfo is None) != (ended.tzinfo is None):
            started = started.replace(tzinfo=None)
            ended = ended.replace(tzinfo=None)
        time_taken_minutes = max(0, round((ended - started).total_seconds() / 60))

    return {
        "score": attempt.score, "correct": attempt.correct_count,
        "wrong": attempt.wrong_count, "unanswered": attempt.unanswered_count,
        "correctScore": test.correct_score if test else 4,
        "wrongPenalty": test.wrong_penalty if test else 1,
        "totalQuestions": total_questions,
        "maxScore": max_score,
        "timeTakenMinutes": time_taken_minutes,
        "fullName": candidate.full_name,
        "testTitle": test.title if test else None,
        "submittedAt": attempt.submitted_at.isoformat() if attempt.submitted_at else None,
        "startedAt": attempt.started_at.isoformat() if attempt.started_at else None,
    }


# --- Admin Screening Endpoints ---

@router.get("/api/admin/screening/tests")
def admin_list_tests(admin: models.User = Depends(require_admin), db: Session = Depends(get_db)):
    tests = db.query(models.ScreeningTest).order_by(models.ScreeningTest.created_at.desc()).all()
    return [serialize_test(t) for t in tests]


@router.post("/api/admin/screening/tests")
def admin_create_test(req: CreateTestRequest,
                       admin: models.User = Depends(require_admin), db: Session = Depends(get_db)):
    test = models.ScreeningTest(
        title=req.title, timer_minutes=req.timerMinutes,
        correct_score=req.correctScore, wrong_penalty=req.wrongPenalty,
    )
    db.add(test)
    db.flush()
    db.add(models.ScreeningSection(
        test_id=test.id, title="Personal Details",
        section_type="personal_data", order=0,
    ))
    db.commit()
    db.refresh(test)
    return serialize_test(test, include_sections=True)


@router.get("/api/admin/screening/tests/{test_id}")
def admin_get_test(test_id: str,
                   admin: models.User = Depends(require_admin), db: Session = Depends(get_db)):
    test = db.get(models.ScreeningTest, test_id)
    if not test:
        raise HTTPException(404, "Test not found")
    return serialize_test(test, include_sections=True)


@router.post("/api/admin/screening/tests/{test_id}/duplicate")
def admin_duplicate_test(test_id: str,
                       admin: models.User = Depends(require_admin), db: Session = Depends(get_db)):
    test = db.get(models.ScreeningTest, test_id)
    if not test:
        raise HTTPException(404, "Test not found")
    
    match = re.search(r' (\d+)$', test.title)
    if match:
        num = int(match.group(1)) + 1
        new_title = re.sub(r' \d+$', f' {num}', test.title)
    else:
        new_title = f"{test.title} 1"
    
    new_test = models.ScreeningTest(
        title=new_title,
        timer_minutes=test.timer_minutes,
        correct_score=test.correct_score,
        wrong_penalty=test.wrong_penalty,
        is_active=test.is_active,
    )
    db.add(new_test)
    db.flush()
    
    for section in test.sections:
        new_section = models.ScreeningSection(
            test_id=new_test.id,
            title=section.title,
            section_type=section.section_type,
            passage=section.passage,
            order=section.order
        )
        db.add(new_section)
        db.flush()
        
        for q in section.questions:
            new_q = models.ScreeningQuestion(
                section_id=new_section.id,
                prompt=q.prompt,
                options=q.options,
                answer=q.answer,
                image_urls=q.image_urls,
                order=q.order
            )
            db.add(new_q)
            
    db.commit()
    db.refresh(new_test)
    return serialize_test(new_test, include_sections=True)


@router.patch("/api/admin/screening/tests/{test_id}")
def admin_update_test(test_id: str, req: CreateTestRequest,
                       admin: models.User = Depends(require_admin), db: Session = Depends(get_db)):
    test = db.get(models.ScreeningTest, test_id)
    if not test:
        raise HTTPException(404, "Test not found")
    test.title = req.title
    test.timer_minutes = req.timerMinutes
    test.correct_score = req.correctScore
    test.wrong_penalty = req.wrongPenalty
    db.commit()
    return serialize_test(test, include_sections=True)


@router.patch("/api/admin/screening/tests/{test_id}/toggle")
def admin_toggle_test(test_id: str,
                       admin: models.User = Depends(require_admin), db: Session = Depends(get_db)):
    test = db.get(models.ScreeningTest, test_id)
    if not test:
        raise HTTPException(404, "Test not found")
    test.is_active = not test.is_active
    db.commit()
    return {"isActive": test.is_active}


@router.delete("/api/admin/screening/tests/{test_id}")
def admin_delete_test(test_id: str,
                       admin: models.User = Depends(require_admin), db: Session = Depends(get_db)):
    test = db.get(models.ScreeningTest, test_id)
    if not test:
        raise HTTPException(404, "Test not found")
    db.delete(test)
    db.commit()
    return {"ok": True}


@router.post("/api/admin/screening/tests/{test_id}/sections")
def admin_add_section(test_id: str, req: SectionRequest,
                       admin: models.User = Depends(require_admin), db: Session = Depends(get_db)):
    test = db.get(models.ScreeningTest, test_id)
    if not test:
        raise HTTPException(404, "Test not found")
    existing_orders = [s.order for s in test.sections]
    next_order = max(existing_orders) + 1 if existing_orders else 1
    sec = models.ScreeningSection(
        test_id=test_id, title=req.title,
        section_type=req.section_type, passage=req.passage,
        order=next_order,
    )
    db.add(sec)
    db.commit()
    db.refresh(test)
    return serialize_test(test, include_sections=True)


@router.patch("/api/admin/screening/tests/{test_id}/sections/{section_id}")
def admin_update_section(test_id: str, section_id: str, req: SectionRequest,
                          admin: models.User = Depends(require_admin), db: Session = Depends(get_db)):
    sec = db.get(models.ScreeningSection, section_id)
    if not sec or sec.test_id != test_id:
        raise HTTPException(404, "Section not found")
    sec.title = req.title
    sec.passage = req.passage
    db.commit()
    test = db.get(models.ScreeningTest, test_id)
    return serialize_test(test, include_sections=True)


@router.delete("/api/admin/screening/tests/{test_id}/sections/{section_id}")
def admin_delete_section(test_id: str, section_id: str,
                          admin: models.User = Depends(require_admin), db: Session = Depends(get_db)):
    sec = db.get(models.ScreeningSection, section_id)
    if not sec or sec.test_id != test_id:
        raise HTTPException(404, "Section not found")
    if sec.section_type == "personal_data":
        raise HTTPException(400, "Cannot delete the personal data section")
    db.delete(sec)
    db.commit()
    test = db.get(models.ScreeningTest, test_id)
    return serialize_test(test, include_sections=True)


@router.put("/api/admin/screening/tests/{test_id}/sections/{section_id}/questions")
def admin_save_section_questions(
    test_id: str, section_id: str, req: BulkQuestionsRequest,
    admin: models.User = Depends(require_admin), db: Session = Depends(get_db),
):
    """Replace all questions in a section."""
    sec = db.get(models.ScreeningSection, section_id)
    if not sec or sec.test_id != test_id:
        raise HTTPException(404, "Section not found")
    db.query(models.ScreeningQuestion).filter_by(section_id=section_id).delete()
    for i, q in enumerate(req.questions):
        db.add(models.ScreeningQuestion(
            section_id=section_id, prompt=q.prompt, options=q.options,
            answer=q.answer, image_urls=q.imageUrls, order=i,
        ))
    db.commit()
    test = db.get(models.ScreeningTest, test_id)
    return serialize_test(test, include_sections=True)


@router.post("/api/admin/screening/upload-question-image")
async def admin_upload_screening_question_image(
    file: UploadFile = File(...),
    admin: models.User = Depends(require_admin),
):
    """Upload an image for a screening MCQ question."""
    from PIL import Image
    await validate_safe_upload(file)
    data = await file.read()
    try:
        img = Image.open(io.BytesIO(data))
        img.verify()
    except Exception:
        raise HTTPException(400, "Invalid image file")
    img2 = Image.open(io.BytesIO(data)).convert("RGB")
    img2.thumbnail((1000, 1000))
    filename = f"{uuid.uuid4().hex}.jpg"
    tmp_path = os.path.join(UPLOAD_DIR, f"_tmp_{filename}")
    img2.save(tmp_path, "JPEG", quality=88)
    storage.save("screening-questions", filename, tmp_path)
    return {"imageUrl": f"/api/uploads/screening-questions/{filename}"}


@router.get("/api/admin/screening/candidates")
def admin_list_candidates(
    test_id: Optional[str] = None,
    admin: models.User = Depends(require_admin), db: Session = Depends(get_db)
):
    q = db.query(models.ScreeningCandidate)
    if test_id:
        q = q.filter_by(test_id=test_id)
    candidates = q.order_by(models.ScreeningCandidate.created_at.desc()).all()
    result = []
    for c in candidates:
        test = db.get(models.ScreeningTest, c.test_id)
        att = c.attempt
        result.append({
            "id": c.id, "fullName": c.full_name, "mobileNumber": c.mobile_number,
            "status": c.status, "isActive": c.is_active,
            "testId": c.test_id, "testTitle": test.title if test else "—",
            "createdAt": c.created_at.isoformat() if c.created_at else None,
            "startedAt": att.started_at.isoformat() if att and att.started_at else None,
            "submittedAt": att.submitted_at.isoformat() if att and att.submitted_at else None,
            "score": att.score if att else None,
            "tabSwitchCount": att.tab_switch_count or 0 if att else 0,
        })
    return result


@router.post("/api/admin/screening/candidates")
def admin_create_candidate(
    req: CreateCandidateRequest,
    admin: models.User = Depends(require_admin), db: Session = Depends(get_db)
):
    test = db.get(models.ScreeningTest, req.testId)
    if not test:
        raise HTTPException(404, "Test not found")
    c = models.ScreeningCandidate(
        test_id=req.testId, full_name=req.fullName.strip(),
        password_hash=hash_password(req.password),
        mobile_number=req.mobileNumber,
    )
    db.add(c)
    db.commit()
    db.refresh(c)
    return {
        "id": c.id, "fullName": c.full_name, "mobileNumber": c.mobile_number,
        "status": c.status, "testId": c.test_id, "testTitle": test.title,
        "createdAt": c.created_at.isoformat() if c.created_at else None,
        "startedAt": None, "submittedAt": None, "score": None,
    }


@router.patch("/api/admin/screening/candidates/{cand_id}")
def admin_update_candidate(
    cand_id: str, req: dict,
    admin: models.User = Depends(require_admin), db: Session = Depends(get_db)
):
    c = db.get(models.ScreeningCandidate, cand_id)
    if not c:
        raise HTTPException(404, "Candidate not found")
    if "isActive" in req:
        c.is_active = req["isActive"]
    if "fullName" in req and req["fullName"].strip():
        c.full_name = req["fullName"].strip()
    if "mobileNumber" in req:
        c.mobile_number = req["mobileNumber"]
    if "password" in req and req["password"]:
        c.password_hash = hash_password(req["password"])
    if "testId" in req and req["testId"] and req["testId"] != c.test_id:
        new_test = db.get(models.ScreeningTest, req["testId"])
        if not new_test:
            raise HTTPException(404, "Test not found")
        if c.attempt:
            db.delete(c.attempt)
        c.test_id = req["testId"]
        c.status = "pending"
    db.commit()
    return {"ok": True}


@router.delete("/api/admin/screening/candidates/{cand_id}")
def admin_delete_candidate(
    cand_id: str,
    admin: models.User = Depends(require_admin), db: Session = Depends(get_db)
):
    c = db.get(models.ScreeningCandidate, cand_id)
    if not c:
        raise HTTPException(404, "Candidate not found")
    db.delete(c)
    db.commit()
    return {"ok": True}


@router.get("/api/admin/screening/results")
def admin_screening_results(
    test_id: Optional[str] = None,
    admin: models.User = Depends(require_admin), db: Session = Depends(get_db)
):
    q = db.query(models.ScreeningAttempt)
    if test_id:
        q = q.filter_by(test_id=test_id)
    attempts = q.order_by(models.ScreeningAttempt.submitted_at.desc()).all()
    rows = []
    for att in attempts:
        c = db.get(models.ScreeningCandidate, att.candidate_id)
        test = db.get(models.ScreeningTest, att.test_id)
        if not c:
            continue
        time_taken = None
        if att.started_at and att.submitted_at:
            delta = att.submitted_at - att.started_at
            time_taken = round(delta.total_seconds() / 60, 1)
        total_q = sum(
            len(s.questions) for s in (test.sections if test else [])
            if s.section_type == "mcq"
        )
        max_score = total_q * (test.correct_score if test else 4)
        rows.append({
            "candidateId": c.id, "fullName": c.full_name,
            "mobileNumber": c.mobile_number or att.personal_data.get("mobile", "") if att.personal_data else (c.mobile_number or ""),
            "testId": att.test_id, "testTitle": test.title if test else "—",
            "startedAt": att.started_at.isoformat() if att.started_at else None,
            "submittedAt": att.submitted_at.isoformat() if att.submitted_at else None,
            "timeTakenMinutes": time_taken,
            "correct": att.correct_count or 0,
            "wrong": att.wrong_count or 0,
            "unanswered": att.unanswered_count or 0,
            "score": att.score or 0,
            "maxScore": max_score,
            "totalQuestions": total_q,
            "personalData": att.personal_data,
            "tabSwitchCount": att.tab_switch_count or 0,
        })
    return rows


@router.get("/api/admin/screening/candidates/{cand_id}/answers")
def admin_get_candidate_answers(
    cand_id: str,
    admin: models.User = Depends(require_admin),
    db: Session = Depends(get_db)
):
    """Retrieve detailed questions and answered options for a candidate's attempt."""
    candidate = db.get(models.ScreeningCandidate, cand_id)
    if not candidate:
        raise HTTPException(404, "Candidate not found")
    attempt = candidate.attempt
    if not attempt:
        raise HTTPException(404, "No attempt found for this candidate")
    test = db.get(models.ScreeningTest, attempt.test_id)
    if not test:
        raise HTTPException(404, "Test not found")

    time_taken = None
    if attempt.started_at and attempt.submitted_at:
        delta = attempt.submitted_at - attempt.started_at
        time_taken = round(delta.total_seconds() / 60, 1)

    total_q = sum(
        len(s.questions) for s in (test.sections or [])
        if s.section_type == "mcq"
    )
    max_score = total_q * test.correct_score

    section_answers = attempt.section_answers or {}
    sections_data = []

    for sec in (test.sections or []):
        if sec.section_type == "personal_data":
            sections_data.append({
                "id": sec.id,
                "title": sec.title,
                "type": sec.section_type,
                "order": sec.order,
                "passage": None,
                "questions": [],
                "questionCount": 0,
                "correctCount": 0,
                "wrongCount": 0,
                "unansweredCount": 0,
                "score": 0,
                "maxScore": 0,
            })
            continue

        raw_answers = section_answers.get(sec.id, [])
        qs_data = []
        for i, q in enumerate(sec.questions):
            chosen = raw_answers[i] if i < len(raw_answers) else None
            is_correct = (chosen is not None) and (chosen == q.answer)
            is_unanswered = (chosen is None)
            is_wrong = (chosen is not None) and (chosen != q.answer)
            points = test.correct_score if is_correct else (-test.wrong_penalty if is_wrong else 0)

            qs_data.append({
                "id": q.id,
                "order": q.order,
                "prompt": q.prompt,
                "options": q.options or [],
                "imageUrls": q.image_urls or [],
                "correctAnswer": q.answer,
                "selectedAnswer": chosen,
                "isCorrect": is_correct,
                "isWrong": is_wrong,
                "isUnanswered": is_unanswered,
                "points": points,
            })

        sec_correct = sum(1 for q in qs_data if q["isCorrect"])
        sec_wrong = sum(1 for q in qs_data if q["isWrong"])
        sec_unanswered = sum(1 for q in qs_data if q["isUnanswered"])
        sec_score = sum(q["points"] for q in qs_data)
        sec_max = len(qs_data) * test.correct_score

        sections_data.append({
            "id": sec.id,
            "title": sec.title,
            "type": sec.section_type,
            "order": sec.order,
            "passage": sec.passage,
            "questions": qs_data,
            "questionCount": len(qs_data),
            "correctCount": sec_correct,
            "wrongCount": sec_wrong,
            "unansweredCount": sec_unanswered,
            "score": sec_score,
            "maxScore": sec_max,
        })

    return {
        "candidate": {
            "id": candidate.id,
            "fullName": candidate.full_name,
            "mobileNumber": candidate.mobile_number or (attempt.personal_data.get("mobile") if attempt.personal_data else None),
            "status": candidate.status,
            "personalData": attempt.personal_data or {},
            "startedAt": attempt.started_at.isoformat() if attempt.started_at else None,
            "submittedAt": attempt.submitted_at.isoformat() if attempt.submitted_at else None,
            "timeTakenMinutes": time_taken,
            "tabSwitchCount": attempt.tab_switch_count or 0,
        },
        "test": {
            "id": test.id,
            "title": test.title,
            "timerMinutes": test.timer_minutes,
            "correctScore": test.correct_score,
            "wrongPenalty": test.wrong_penalty,
            "totalQuestions": total_q,
            "maxScore": max_score,
        },
        "score": attempt.score or 0,
        "maxScore": max_score,
        "correctCount": attempt.correct_count or 0,
        "wrongCount": attempt.wrong_count or 0,
        "unansweredCount": attempt.unanswered_count or 0,
        "accuracy": round((attempt.correct_count / (attempt.correct_count + attempt.wrong_count)) * 100, 1) if (attempt.correct_count is not None and attempt.wrong_count is not None and (attempt.correct_count + attempt.wrong_count) > 0) else None,
        "sections": sections_data,
    }


@router.get("/api/admin/screening/results.xlsx")
def admin_screening_results_xlsx(
    test_id: Optional[str] = None,
    admin: models.User = Depends(require_admin), db: Session = Depends(get_db)
):
    """Download screening results as a styled Excel workbook."""
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter

    rows_data = admin_screening_results(test_id=test_id, admin=admin, db=db)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Screening Results"

    header_fill = PatternFill("solid", fgColor="1A2744")
    header_font = Font(bold=True, color="FFFFFF", size=11)
    center = Alignment(horizontal="center", vertical="center")
    thin = Border(
        left=Side(style="thin"), right=Side(style="thin"),
        top=Side(style="thin"), bottom=Side(style="thin"),
    )

    headers = [
        "Name", "Mobile Number", "Test", "Start Time", "Submit Time",
        "Time Taken (min)", "Correct", "Wrong", "Unanswered", "Score",
        "Remark", "Shortlisted",
    ]
    ws.append(headers)
    for col_idx, _ in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = center
        cell.border = thin
    ws.row_dimensions[1].height = 22

    alt_fill = PatternFill("solid", fgColor="EEF2FF")
    for row_i, r in enumerate(rows_data, 2):
        def fmt_dt(iso):
            if not iso:
                return ""
            try:
                from datetime import datetime as dt
                return dt.fromisoformat(iso.replace("Z", "+00:00")).strftime("%d %b %Y %H:%M")
            except Exception:
                return iso

        ws.append([
            r["fullName"], r["mobileNumber"] or "", r["testTitle"],
            fmt_dt(r["startedAt"]), fmt_dt(r["submittedAt"]),
            r["timeTakenMinutes"] or "",
            r["correct"], r["wrong"], r["unanswered"],
            r["score"],
            "",  # Remark
            "",  # Shortlisted
        ])
        fill = alt_fill if row_i % 2 == 0 else None
        for col_idx in range(1, len(headers) + 1):
            cell = ws.cell(row=row_i, column=col_idx)
            cell.border = thin
            cell.alignment = Alignment(vertical="center")
            if fill:
                cell.fill = fill

    col_widths = [28, 18, 28, 20, 20, 18, 10, 10, 12, 10, 15, 30]
    for i, w in enumerate(col_widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=screening-results.xlsx"},
    )
