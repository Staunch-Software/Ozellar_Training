"""Orientation program router: Admin program & enrollment management, crew task completion, vessel approver verification."""
import io
import os
import tempfile
import uuid
from datetime import datetime, timezone
from zoneinfo import ZoneInfo
from typing import Optional

IST = ZoneInfo("Asia/Kolkata")


def to_ist_iso(dt: Optional[datetime]) -> Optional[str]:
    """Serializes datetime into unambiguous ISO 8601 string with +05:30 IST offset."""
    if not dt:
        return None
    if dt.tzinfo is None:
        return dt.isoformat() + "+05:30"
    return dt.astimezone(IST).isoformat()


from fastapi import APIRouter, Depends, HTTPException, Form, File, UploadFile
from fastapi.responses import StreamingResponse
from pydantic import BaseModel
from sqlalchemy.orm import Session
from sqlalchemy import func

from app.database import get_db
from app import models, orientation_ranks, storage
from app.auth import (
    get_current_user, require_admin, require_super_admin,
    require_vessel_approver, normalize_name
)
from app.security_upload import validate_safe_upload
from app.notifications_service import notify

router = APIRouter(tags=["Orientation"])

CANDIDATES_CAP = 2000


def orientation_program_summary(program):
    return {
        "id": program.id, "title": program.title, "subtitle": program.subtitle,
        "department": program.department, "fromRank": program.from_rank,
        "toRank": program.to_rank, "isActive": bool(program.is_active),
        "order": program.order, "taskCount": len(program.tasks),
        "enrollmentCount": len(program.enrollments),
    }


def orientation_task_detail(t):
    return {
        "id": t.id, "title": t.title, "description": t.description,
        "order": t.order, "requiresProof": bool(t.requires_proof),
    }


def _next_orientation_task_order(program) -> int:
    return (max((t.order for t in program.tasks), default=-1)) + 1


class OrientationProgramRequest(BaseModel):
    title: str
    subtitle: str | None = None
    department: str            # 'deck' | 'engine'
    fromRank: str | None = None
    toRank: str | None = None


class OrientationTaskRequest(BaseModel):
    title: str
    description: str | None = None
    requiresProof: bool = False


class OrientationReorderRequest(BaseModel):
    order: list[str]


class OrientationEnrollRequest(BaseModel):
    learnerId: str
    programId: str
    vesselName: str | None = None
    masterName: str | None = None


class SubmitMultipleTasksRequest(BaseModel):
    task_ids: list[str]


class OrientationVerifyTaskRequest(BaseModel):
    action: str  # 'approve' | 'reject'
    note: str | None = None


class OrientationDecideRequest(BaseModel):
    action: str


class AdminOrientationDecideRequest(BaseModel):
    action: str  # 'approve' | 'reject'


# --- Admin Orientation Programs ---

@router.get("/api/admin/orientation/programs")
def admin_list_orientation_programs(admin: models.User = Depends(require_admin), db: Session = Depends(get_db)):
    programs = db.query(models.OrientationProgram).order_by(models.OrientationProgram.order).all()
    return [orientation_program_summary(p) for p in programs]


@router.post("/api/admin/orientation/programs")
def admin_create_orientation_program(req: OrientationProgramRequest,
                                     admin: models.User = Depends(require_admin), db: Session = Depends(get_db)):
    title = (req.title or "").strip()
    if not title:
        raise HTTPException(400, "Title is required")
    if req.department not in ("deck", "engine"):
        raise HTTPException(400, "Department must be 'deck' or 'engine'")
    max_order = db.query(func.max(models.OrientationProgram.order)).scalar()
    program = models.OrientationProgram(
        title=title, subtitle=(req.subtitle or "").strip() or None,
        department=req.department, from_rank=(req.fromRank or "").strip() or None,
        to_rank=(req.toRank or "").strip() or None,
        order=(max_order + 1) if max_order is not None else 0,
    )
    db.add(program)
    db.commit()
    db.refresh(program)
    return orientation_program_summary(program)


@router.get("/api/admin/orientation/programs/enrollments", include_in_schema=False)
def admin_programs_enrollments_alias(
    program_id: str | None = None,
    status: str | None = None,
    admin: models.User = Depends(require_admin),
    db: Session = Depends(get_db),
):
    """Fallback alias if enrollments is requested under /programs/enrollments."""
    return admin_list_orientation_enrollments(program_id=program_id, status=status, admin=admin, db=db)


@router.get("/api/admin/orientation/programs/results", include_in_schema=False)
def admin_programs_results_alias(
    program_id: str | None = None,
    admin: models.User = Depends(require_admin),
    db: Session = Depends(get_db),
):
    """Fallback alias if results is requested under /programs/results."""
    return admin_orientation_results(program_id=program_id, admin=admin, db=db)


@router.get("/api/admin/orientation/programs/{program_id}")
def admin_get_orientation_program(program_id: str, admin: models.User = Depends(require_admin),
                                  db: Session = Depends(get_db)):
    program = db.get(models.OrientationProgram, program_id)
    if not program:
        raise HTTPException(404, "Program not found")
    return {
        **orientation_program_summary(program),
        "tasks": [orientation_task_detail(t) for t in sorted(program.tasks, key=lambda t: t.order)],
    }


@router.put("/api/admin/orientation/programs/{program_id}")
def admin_update_orientation_program(program_id: str, req: OrientationProgramRequest,
                                     admin: models.User = Depends(require_admin), db: Session = Depends(get_db)):
    program = db.get(models.OrientationProgram, program_id)
    if not program:
        raise HTTPException(404, "Program not found")
    title = (req.title or "").strip()
    if not title:
        raise HTTPException(400, "Title is required")
    if req.department not in ("deck", "engine"):
        raise HTTPException(400, "Department must be 'deck' or 'engine'")
    program.title = title
    program.subtitle = (req.subtitle or "").strip() or None
    program.department = req.department
    program.from_rank = (req.fromRank or "").strip() or None
    program.to_rank = (req.toRank or "").strip() or None
    db.commit()
    return orientation_program_summary(program)


@router.patch("/api/admin/orientation/programs/{program_id}/toggle")
def admin_toggle_orientation_program(program_id: str, admin: models.User = Depends(require_admin),
                                     db: Session = Depends(get_db)):
    program = db.get(models.OrientationProgram, program_id)
    if not program:
        raise HTTPException(404, "Program not found")
    program.is_active = not program.is_active
    db.commit()
    return {"isActive": program.is_active}


@router.delete("/api/admin/orientation/programs/{program_id}")
def admin_delete_orientation_program(program_id: str, admin: models.User = Depends(require_admin),
                                     db: Session = Depends(get_db)):
    program = db.get(models.OrientationProgram, program_id)
    if not program:
        raise HTTPException(404, "Program not found")
    db.delete(program)
    db.commit()
    return {"ok": True}


@router.get("/api/admin/orientation/programs.xlsx")
def admin_orientation_programs_xlsx(admin: models.User = Depends(require_admin), db: Session = Depends(get_db)):
    """Export all orientation programs as an Excel workbook — one sheet per program."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment, Border, Side

    programs = db.query(models.OrientationProgram).order_by(models.OrientationProgram.order).all()
    wb = Workbook()
    wb.remove(wb.active)

    thin = Side(style="thin")
    border_bottom = Border(bottom=thin)

    for prog in programs:
        sheet_name = (prog.title or "Program")[:31].replace("/", "-").replace("\\", "-").replace("?", "").replace("*", "").replace("[", "").replace("]", "")
        ws = wb.create_sheet(title=sheet_name)
        ws.column_dimensions["A"].width = 120

        dept_label = "Engine" if prog.department == "engine" else "Deck"
        approver_role = "Chief Engineer" if prog.department == "engine" else "Master"

        ws.append([prog.title or ""])
        title_cell = ws["A1"]
        title_cell.font = Font(bold=True, size=14)
        title_cell.alignment = Alignment(wrap_text=True)

        subtitle = prog.subtitle or f"{dept_label} Mentoring Program"
        ws.append([subtitle])
        ws["A2"].font = Font(bold=True, size=11)

        ws.append([""])

        rank_from = prog.from_rank or ""
        rank_to = prog.to_rank or ""
        ws.append([f"{'Name':<54}{'ID':<40}{'Rank':<20}"])
        ws["A4"].font = Font(bold=True)
        ws["A4"].border = border_bottom

        ws.append([""])

        row_num = 6
        tasks = sorted(prog.tasks, key=lambda t: t.order)
        for task in tasks:
            task_name_line = f"Task Name                           {task.title}"
            ws.append([task_name_line])
            tc = ws.cell(row=row_num, column=1)
            tc.font = Font(bold=True)
            tc.alignment = Alignment(wrap_text=True)
            row_num += 1

            desc_text = (task.description or "").strip()
            proof_note = "\n(Relevant task supporting documents to be presented during Mentoring Review at FMTI/office)" if task.requires_proof else ""
            full_desc = f"Task Description                 {desc_text}{proof_note}"
            ws.append([full_desc])
            dc = ws.cell(row=row_num, column=1)
            dc.alignment = Alignment(wrap_text=True, vertical="top")
            dc.font = Font(size=10)
            line_count = max(1, full_desc.count("\n") + 1)
            ws.row_dimensions[row_num].height = max(15, line_count * 14)
            row_num += 1

            ws.append([f"{approver_role}'s Initials:                                             Date :"])
            init_cell = ws.cell(row=row_num, column=1)
            init_cell.font = Font(bold=True, size=10)
            init_cell.border = border_bottom
            row_num += 1

            ws.append([""])
            row_num += 1

        ws.row_dimensions[1].height = 22
        ws.row_dimensions[2].height = 18

    if not wb.sheetnames:
        wb.create_sheet("No Programs")

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=orientation-programs.xlsx"},
    )


@router.post("/api/admin/orientation/programs/{program_id}/tasks")
def admin_add_orientation_task(program_id: str, req: OrientationTaskRequest,
                               admin: models.User = Depends(require_admin), db: Session = Depends(get_db)):
    program = db.get(models.OrientationProgram, program_id)
    if not program:
        raise HTTPException(404, "Program not found")
    title = (req.title or "").strip()
    if not title:
        raise HTTPException(400, "Title is required")
    task = models.OrientationTask(
        program_id=program_id, title=title, description=(req.description or "").strip() or None,
        order=_next_orientation_task_order(program),
        requires_proof=bool(req.requiresProof),
    )
    db.add(task)
    db.commit()
    db.refresh(task)
    return orientation_task_detail(task)


@router.put("/api/admin/orientation/programs/{program_id}/tasks/{task_id}")
def admin_update_orientation_task(program_id: str, task_id: str, req: OrientationTaskRequest,
                                  admin: models.User = Depends(require_admin), db: Session = Depends(get_db)):
    task = db.get(models.OrientationTask, task_id)
    if not task or task.program_id != program_id:
        raise HTTPException(404, "Task not found")
    title = (req.title or "").strip()
    if not title:
        raise HTTPException(400, "Title is required")
    task.title = title
    task.description = (req.description or "").strip() or None
    task.requires_proof = bool(req.requiresProof)
    db.commit()
    return orientation_task_detail(task)


@router.delete("/api/admin/orientation/programs/{program_id}/tasks/{task_id}")
def admin_delete_orientation_task(program_id: str, task_id: str,
                                  admin: models.User = Depends(require_admin), db: Session = Depends(get_db)):
    task = db.get(models.OrientationTask, task_id)
    if not task or task.program_id != program_id:
        raise HTTPException(404, "Task not found")
    db.delete(task)
    db.flush()
    program = db.get(models.OrientationProgram, program_id)
    if program:
        remaining = sorted(program.tasks, key=lambda t: t.order)
        for i, r in enumerate(remaining):
            r.order = i
    db.commit()
    return {"ok": True}


@router.put("/api/admin/orientation/programs/{program_id}/reorder")
def admin_reorder_orientation_tasks(program_id: str, req: OrientationReorderRequest,
                                    admin: models.User = Depends(require_admin), db: Session = Depends(get_db)):
    program = db.get(models.OrientationProgram, program_id)
    if not program:
        raise HTTPException(404, "Program not found")
    if set(req.order) != {t.id for t in program.tasks}:
        raise HTTPException(400, "Order must include exactly the program's current tasks")
    by_id = {t.id: t for t in program.tasks}
    for i, tid in enumerate(req.order):
        by_id[tid].order = i
    db.commit()
    return {"ok": True}


# --- Candidates & Enrollments ---

def orientation_candidate_view(db, u, program_id=None):
    q = db.query(models.OrientationEnrollment).filter_by(learner_id=u.id)
    if program_id:
        q = q.filter_by(program_id=program_id)
    enrollment = q.order_by(models.OrientationEnrollment.created_at.desc()).first()
    return {
        "id": u.id, "name": u.full_name, "rank": u.rank, "vessel": u.current_vessel,
        "empStatus": u.emp_status, "nationality": u.nationality,
        "department": orientation_ranks.department_for_rank(u.rank),
        "enrollment": ({
            "id": enrollment.id, "programId": enrollment.program_id,
            "programTitle": enrollment.program.title if enrollment.program else None,
            "status": enrollment.status,
            "vesselName": enrollment.vessel_name,
            "masterName": enrollment.master_name,
        } if enrollment else None),
    }


@router.get("/api/admin/orientation/vessels")
def admin_list_orientation_vessels(admin: models.User = Depends(require_admin), db: Session = Depends(get_db)):
    _norm = orientation_ranks._norm
    learners = db.query(models.User).filter(
        models.User.role == "learner",
        models.User.is_active == True,
        models.User.current_vessel != None,
        models.User.emp_status == "SAIL",
    ).all()
    vessels: dict[str, dict] = {}
    for u in learners:
        v = (u.current_vessel or "").strip()
        if not v:
            continue
        if v not in vessels:
            vessels[v] = {"vessel": v, "masters": [], "chiefEngineers": [], "crewCount": 0}
        vessels[v]["crewCount"] += 1
        r = _norm(u.rank)
        candidate = {
            "id": u.id, "name": u.full_name,
            "signOnDate": u.sign_on_date.isoformat() if u.sign_on_date else None,
            "reliefDate": u.relief_date.isoformat() if u.relief_date else None,
        }
        if r == "MASTER":
            vessels[v]["masters"].append(candidate)
        elif r == "CHIEF ENGINEER":
            vessels[v]["chiefEngineers"].append(candidate)

    def annotate(candidates: list) -> list:
        if len(candidates) < 2:
            return candidates
        dated = [c for c in candidates if c["reliefDate"]]
        if len(dated) < 2 or len({c["reliefDate"] for c in dated}) < 2:
            return candidates
        soonest = min(dated, key=lambda c: c["reliefDate"])
        soonest["signingOffSoon"] = True
        latest = max(dated, key=lambda c: c["reliefDate"])
        latest["defaultPick"] = True
        return candidates

    result = []
    for v in sorted(vessels.values(), key=lambda x: x["vessel"]):
        v["masters"] = annotate(v["masters"])
        v["chiefEngineers"] = annotate(v["chiefEngineers"])
        v["master"] = v["masters"][0]["name"] if len(v["masters"]) == 1 else None
        v["chiefEngineer"] = v["chiefEngineers"][0]["name"] if len(v["chiefEngineers"]) == 1 else None
        result.append(v)
    return result


@router.get("/api/admin/orientation/candidates")
def admin_list_orientation_candidates(
    q: str = "",
    program_id: str | None = None,
    rank: str | None = None,
    vessel: str | None = None,
    on_sail: bool = False,
    admin: models.User = Depends(require_admin),
    db: Session = Depends(get_db),
):
    users = db.query(models.User).filter_by(role="learner", is_active=True).all()
    eligible = [u for u in users if orientation_ranks.is_eligible_crew(u.rank)]
    if q:
        query = normalize_name(q)
        eligible = [u for u in eligible if query in normalize_name(u.full_name)]
    if rank:
        rank_upper = rank.strip().upper()
        eligible = [u for u in eligible if (u.rank or "").strip().upper() == rank_upper]
    if vessel:
        vessel_upper = vessel.strip().upper()
        eligible = [u for u in eligible if (u.current_vessel or "").strip().upper() == vessel_upper]
    if on_sail:
        eligible = [u for u in eligible if (u.emp_status or "").strip().upper() == "SAIL"]
    eligible.sort(key=lambda u: u.full_name)
    return [orientation_candidate_view(db, u, program_id) for u in eligible[:CANDIDATES_CAP]]


@router.get("/api/admin/orientation/enrollments")
def admin_list_orientation_enrollments(
    program_id: str | None = None,
    status: str | None = None,
    admin: models.User = Depends(require_admin),
    db: Session = Depends(get_db),
):
    q = db.query(models.OrientationEnrollment)
    if program_id:
        q = q.filter_by(program_id=program_id)
    if status:
        q = q.filter_by(status=status)
    enrollments = q.order_by(models.OrientationEnrollment.created_at.desc()).all()
    result = []
    for e in enrollments:
        learner = db.get(models.User, e.learner_id)
        total = len(e.program.tasks) if e.program else 0
        done = sum(1 for c in e.completions if c.is_completed)
        pct = round(done / total * 100) if total else 0
        result.append({
            "id": e.id,
            "learnerId": e.learner_id,
            "learnerName": learner.full_name if learner else None,
            "rank": learner.rank if learner else None,
            "vessel": e.vessel_name or (learner.current_vessel if learner else None),
            "empStatus": learner.emp_status if learner else None,
            "department": orientation_ranks.department_for_rank(learner.rank) if learner else None,
            "programId": e.program_id,
            "programTitle": e.program.title if e.program else None,
            "status": e.status,
            "vesselName": e.vessel_name,
            "masterName": e.master_name,
            "completedCount": done,
            "totalCount": total,
            "progressPct": pct,
            "createdAt": to_ist_iso(e.created_at),
        })
    return result


@router.post("/api/admin/orientation/enrollments")
def admin_create_orientation_enrollment(req: OrientationEnrollRequest,
                                        admin: models.User = Depends(require_admin), db: Session = Depends(get_db)):
    learner = db.get(models.User, req.learnerId)
    if not learner or learner.role != "learner":
        raise HTTPException(404, "Crew member not found")
    if not orientation_ranks.is_eligible_crew(learner.rank):
        raise HTTPException(400, "This crew member's rank is not eligible for Orientation Program")
    program = db.get(models.OrientationProgram, req.programId)
    if not program:
        raise HTTPException(404, "Program not found")
    existing = db.query(models.OrientationEnrollment).filter_by(
        learner_id=learner.id, program_id=program.id).first()
    if existing and existing.status in ("in_progress", "submitted"):
        raise HTTPException(400, "This crew member is already enrolled in this program")
    vessel_name = (req.vesselName or "").strip() or learner.current_vessel
    enrollment = models.OrientationEnrollment(
        program_id=program.id, learner_id=learner.id, assigned_by=admin.id, status="in_progress",
        vessel_name=vessel_name,
        master_name=(req.masterName or "").strip() or None,
    )
    db.add(enrollment)
    db.flush()
    for task in program.tasks:
        db.add(models.OrientationTaskCompletion(enrollment_id=enrollment.id, task_id=task.id))
    db.commit()
    return {"ok": True, "enrollmentId": enrollment.id}


@router.delete("/api/admin/orientation/enrollments/{enrollment_id}")
def admin_delete_orientation_enrollment(enrollment_id: str,
                                        admin: models.User = Depends(require_admin), db: Session = Depends(get_db)):
    enrollment = db.get(models.OrientationEnrollment, enrollment_id)
    if not enrollment:
        raise HTTPException(404, "Enrollment not found")
    db.delete(enrollment)
    db.commit()
    return {"ok": True}


# --- Results / Monitoring ---

def orientation_enrollment_result(db, e):
    learner = db.get(models.User, e.learner_id)
    tasks = e.program.tasks if e.program else []
    total = len(tasks)
    completions_map = {c.task_id: c for c in e.completions}
    done = sum(1 for t in tasks if completions_map.get(t.id) and completions_map[t.id].is_completed)
    pct = round(done / total * 100) if total else 0

    submitted_completions = [c for c in e.completions if c.status in ("pending_review", "approved", "rejected") and c.completed_at]
    submitted_at = min((c.completed_at for c in submitted_completions), default=None)

    decided_at = None
    if e.status in ("approved", "rejected"):
        verified_times = [c.verified_at for c in e.completions if c.verified_at]
        decided_at = max(verified_times, default=None)

    submission = (db.query(models.OrientationSubmission)
                  .filter_by(enrollment_id=e.id)
                  .order_by(models.OrientationSubmission.submitted_at.desc()).first())
    if submission:
        if not submitted_at and submission.submitted_at:
            submitted_at = submission.submitted_at
        if not decided_at and submission.decided_at:
            decided_at = submission.decided_at

    decided_by_user = db.get(models.User, submission.decided_by) if submission and submission.decided_by else None
    department = e.program.department if e.program else None
    master_label = "Chief Engineer" if department == "engine" else "Master"
    return {
        "enrollmentId": e.id,
        "learnerName": learner.full_name if learner else None,
        "rank": learner.rank if learner else None,
        "vessel": (e.vessel_name or (learner.current_vessel if learner else None)),
        "programId": e.program_id,
        "programTitle": e.program.title if e.program else None,
        "status": e.status,
        "completedCount": done, "totalCount": total, "progressPct": pct,
        "submittedAt": to_ist_iso(submitted_at),
        "decidedAt": to_ist_iso(decided_at),
        "masterLabel": master_label,
        "masterName": e.master_name,
        "approvedByName": decided_by_user.full_name if decided_by_user else None,
    }


@router.get("/api/admin/orientation/results")
def admin_orientation_results(program_id: str | None = None,
                              admin: models.User = Depends(require_admin), db: Session = Depends(get_db)):
    q = db.query(models.OrientationEnrollment)
    if program_id:
        q = q.filter_by(program_id=program_id)
    enrollments = q.order_by(models.OrientationEnrollment.created_at.desc()).all()
    return [orientation_enrollment_result(db, e) for e in enrollments]


@router.get("/api/admin/orientation/results.xlsx")
def admin_orientation_results_xlsx(program_id: str | None = None,
                                   admin: models.User = Depends(require_admin), db: Session = Depends(get_db)):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    rows = admin_orientation_results(program_id=program_id, admin=admin, db=db)

    wb = Workbook()
    ws = wb.active
    ws.title = "Orientation Results"

    columns = [
        ("Crew Name", 26), ("Rank", 18), ("Vessel", 22), ("Program", 30),
        ("Status", 14), ("Completed", 11), ("Total", 9), ("Progress %", 12),
        ("Master / CE", 22), ("Approved By", 22), ("Submitted", 18), ("Decided", 18),
    ]
    headers = [c[0] for c in columns]
    ws.append(headers)

    header_fill = PatternFill(start_color="1E3A5F", end_color="1E3A5F", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True, size=11)
    thin = Side(style="thin", color="D9D9D9")
    for col_idx, _ in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = Border(bottom=thin)
    ws.row_dimensions[1].height = 22

    status_fill = {
        "approved": PatternFill(start_color="DCFCE7", end_color="DCFCE7", fill_type="solid"),
        "submitted": PatternFill(start_color="DBEAFE", end_color="DBEAFE", fill_type="solid"),
        "rejected": PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid"),
        "in_progress": PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid"),
    }
    status_label = {
        "approved": "Approved", "submitted": "Awaiting review",
        "rejected": "Rejected", "in_progress": "In progress",
    }
    band_fill = PatternFill(start_color="F7F9FC", end_color="F7F9FC", fill_type="solid")

    for i, r in enumerate(rows):
        row_idx = i + 2
        ws.append([
            r["learnerName"], r["rank"], r["vessel"], r["programTitle"],
            status_label.get(r["status"], r["status"]), r["completedCount"], r["totalCount"],
            r["progressPct"] / 100,
            f'{r["masterLabel"]}: {r["masterName"]}' if r["masterName"] else "—",
            r["approvedByName"] or "—",
            r["submittedAt"][:10] if r["submittedAt"] else "",
            r["decidedAt"][:10] if r["decidedAt"] else "",
        ])
        if i % 2 == 1:
            for col_idx in range(1, len(headers) + 1):
                ws.cell(row=row_idx, column=col_idx).fill = band_fill
        status_cell = ws.cell(row=row_idx, column=5)
        fill = status_fill.get(r["status"])
        if fill:
            status_cell.fill = fill
        status_cell.alignment = Alignment(horizontal="center")
        ws.cell(row=row_idx, column=8).number_format = "0%"
        for col_idx in (6, 7, 8):
            ws.cell(row=row_idx, column=col_idx).alignment = Alignment(horizontal="center")

    for i, (_, w) in enumerate(columns, 1):
        ws.column_dimensions[get_column_letter(i)].width = w

    ws.freeze_panes = "A2"
    last_row = max(len(rows) + 1, 2)
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{last_row}"

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return StreamingResponse(
        buf, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=orientation-results.xlsx"},
    )


# --- Candidate-facing Endpoints ---

def orientation_enrollment_detail(e):
    completions = {c.task_id: c for c in e.completions}
    tasks = sorted(e.program.tasks, key=lambda t: t.order) if e.program else []
    total = len(tasks)
    done = sum(1 for t in tasks if completions.get(t.id) and completions[t.id].is_completed)
    department = e.program.department if e.program else None
    return {
        "id": e.id, "programId": e.program_id,
        "programTitle": e.program.title if e.program else None,
        "programSubtitle": e.program.subtitle if e.program else None,
        "status": e.status,
        "vessel": e.vessel_name,
        "masterName": e.master_name,
        "masterLabel": "Chief Engineer" if department == "engine" else "Master",
        "progressPct": round(done / total * 100) if total else 0,
        "completedCount": done, "totalCount": total,
        "tasks": [{
            "id": t.id, "title": t.title, "description": t.description, "order": t.order,
            "requiresProof": bool(t.requires_proof),
            "isCompleted": bool(completions.get(t.id) and (completions[t.id].is_completed or completions[t.id].status in ("approved", "pending_review"))),
            "note": completions[t.id].note if completions.get(t.id) else None,
            "proofUrls": (completions[t.id].proof_paths or []) if completions.get(t.id) else [],
            "completedAt": (completions[t.id].completed_at.isoformat()
                           if completions.get(t.id) and completions[t.id].completed_at else None),
            "status": completions[t.id].status if completions.get(t.id) else "draft",
            "rejectionNote": getattr(completions[t.id], "rejection_note", None) if completions.get(t.id) else None,
        } for t in tasks],
    }


@router.get("/api/orientation/my-enrollment")
def get_my_orientation_enrollment(user: models.User = Depends(get_current_user), db: Session = Depends(get_db)):
    enrollment = (db.query(models.OrientationEnrollment)
                  .filter(models.OrientationEnrollment.learner_id == user.id,
                         models.OrientationEnrollment.status.in_(["in_progress", "submitted", "rejected", "approved", "master_approved"]))
                  .order_by(models.OrientationEnrollment.created_at.desc()).first())
    if not enrollment:
        return None
    return orientation_enrollment_detail(enrollment)


@router.post("/api/orientation/tasks/{task_id}/complete")
async def complete_orientation_task(task_id: str, completed: bool = Form(...),
                                    note: str | None = Form(None),
                                    files: list[UploadFile] = File(default=[]),
                                    user: models.User = Depends(get_current_user),
                                    db: Session = Depends(get_db)):
    task = db.get(models.OrientationTask, task_id)
    if not task:
        raise HTTPException(404, "Task not found")
    enrollment = (db.query(models.OrientationEnrollment)
                  .filter_by(learner_id=user.id, program_id=task.program_id).first())
    if not enrollment or enrollment.status not in ("in_progress", "rejected"):
        raise HTTPException(403, "You are not actively enrolled in this program")
    completion = db.query(models.OrientationTaskCompletion).filter_by(
        enrollment_id=enrollment.id, task_id=task_id).first()
    if not completion:
        completion = models.OrientationTaskCompletion(enrollment_id=enrollment.id, task_id=task_id, proof_paths=[])
        db.add(completion)
        db.flush()
    completion.is_completed = completed
    completion.completed_at = datetime.now(timezone.utc) if completed else None
    if note is not None:
        completion.note = note
    new_urls = []
    for file in files:
        if not file or not file.filename:
            continue
        await validate_safe_upload(file)
        ext = os.path.splitext(file.filename)[1] or ".bin"
        filename = f"{completion.id}-{uuid.uuid4().hex[:8]}{ext}"
        tmp_path = os.path.join(tempfile.gettempdir(), f"orientation-{filename}")
        with open(tmp_path, "wb") as f:
            f.write(await file.read())
        storage.save("orientation_proofs", filename, tmp_path)
        new_urls.append(f"/api/uploads/orientation_proofs/{filename}")
    if new_urls:
        completion.proof_paths = [*(completion.proof_paths or []), *new_urls]
    db.commit()
    return {"ok": True, "isCompleted": completion.is_completed, "proofUrls": completion.proof_paths or []}


@router.post("/api/orientation/tasks/submit-multiple")
def submit_multiple_orientation_tasks(req: SubmitMultipleTasksRequest, user: models.User = Depends(get_current_user), db: Session = Depends(get_db)):
    if not req.task_ids:
        return {"ok": True, "count": 0}
        
    tasks = db.query(models.OrientationTask).filter(models.OrientationTask.id.in_(req.task_ids)).all()
    if not tasks:
        raise HTTPException(404, "Tasks not found")
        
    program_id = tasks[0].program_id
    enrollment = (db.query(models.OrientationEnrollment)
                  .filter_by(learner_id=user.id, program_id=program_id).first())
    if not enrollment or enrollment.status not in ("in_progress", "rejected", "submitted"):
        raise HTTPException(403, "Not actively enrolled")
        
    completions = db.query(models.OrientationTaskCompletion).filter(
        models.OrientationTaskCompletion.enrollment_id == enrollment.id,
        models.OrientationTaskCompletion.task_id.in_(req.task_ids)
    ).all()
    
    completion_map = {c.task_id: c for c in completions}
    
    count = 0
    for task in tasks:
        completion = completion_map.get(task.id)
        if not completion or not completion.is_completed:
            continue
        if task.requires_proof and not completion.proof_paths:
            continue
            
        completion.status = "pending_review"
        completion.rejection_note = None
        count += 1
        
    if count > 0:
        department = enrollment.program.department
        program_title = enrollment.program.title if enrollment.program else "Orientation Program"
        task_word = "tasks" if count > 1 else "task"
        
        for appr in find_vessel_approvers(db, user.current_vessel, department):
            notify(db, appr.id, "orientation_submitted", f"{count} {task_word} submitted for review",
                   body=f"{user.full_name} submitted {count} {task_word} in {program_title} for your review.",
                   link="/approvals")
        db.commit()
        
    return {"ok": True, "count": count}


@router.post("/api/orientation/tasks/{task_id}/submit")
def submit_orientation_task(task_id: str, user: models.User = Depends(get_current_user), db: Session = Depends(get_db)):
    task = db.get(models.OrientationTask, task_id)
    if not task:
        raise HTTPException(404, "Task not found")
    enrollment = (db.query(models.OrientationEnrollment)
                  .filter_by(learner_id=user.id, program_id=task.program_id).first())
    if not enrollment or enrollment.status not in ("in_progress", "rejected", "submitted"):
        raise HTTPException(403, "Not actively enrolled")
    
    completion = db.query(models.OrientationTaskCompletion).filter_by(
        enrollment_id=enrollment.id, task_id=task_id).first()
    
    if not completion or not completion.is_completed:
        raise HTTPException(400, "Task must be marked completed first")
    if task.requires_proof and not completion.proof_paths:
        raise HTTPException(400, "Attach a document/photo before submitting")
    
    completion.status = "pending_review"
    completion.rejection_note = None
    
    department = enrollment.program.department
    program_title = enrollment.program.title if enrollment.program else "Orientation Program"
    
    for appr in find_vessel_approvers(db, user.current_vessel, department):
        notify(db, appr.id, "orientation_submitted", "Task submitted for review",
               body=f"{user.full_name} submitted a task in {program_title} for your review.",
               link="/approvals")
    db.commit()
    return {"ok": True, "status": completion.status}


@router.delete("/api/orientation/tasks/{task_id}/attachments")
def delete_orientation_task_attachment(task_id: str, url: str,
                                       user: models.User = Depends(get_current_user),
                                       db: Session = Depends(get_db)):
    task = db.get(models.OrientationTask, task_id)
    if not task:
        raise HTTPException(404, "Task not found")
    enrollment = (db.query(models.OrientationEnrollment)
                  .filter_by(learner_id=user.id, program_id=task.program_id).first())
    if not enrollment or enrollment.status not in ("in_progress", "rejected"):
        raise HTTPException(403, "You are not actively enrolled in this program")
    completion = db.query(models.OrientationTaskCompletion).filter_by(
        enrollment_id=enrollment.id, task_id=task_id).first()
    if not completion:
        raise HTTPException(404, "Nothing recorded for this task yet")
    completion.proof_paths = [u for u in (completion.proof_paths or []) if u != url]
    db.commit()
    return {"ok": True, "proofUrls": completion.proof_paths or []}


def find_vessel_approvers(db, vessel: str, department: str):
    if not vessel:
        return []
    candidates = db.query(models.User).filter_by(
        role="learner", is_active=True, current_vessel=vessel).all()
    return [u for u in candidates
            if (info := orientation_ranks.vessel_approver_info(u)) and info["department"] == department]


# --- Approver-facing Endpoints ---

def orientation_submission_detail(db, enrollment):
    learner = db.get(models.User, enrollment.learner_id) if enrollment else None
    tasks = sorted(enrollment.program.tasks, key=lambda t: t.order) if enrollment and enrollment.program else []
    completions_map = {c.task_id: c for c in enrollment.completions} if enrollment else {}
    
    tasks_out = []
    for t in tasks:
        c = completions_map.get(t.id)
        tasks_out.append({
            "id": t.id,
            "title": t.title,
            "description": t.description,
            "requiresProof": t.requires_proof,
            "isCompleted": c.is_completed if c else False,
            "proofUrls": c.proof_paths if c else [],
            "note": c.note if c else None,
            "completedAt": to_ist_iso(c.completed_at) if c and c.completed_at else None,
            "status": c.status if c else "draft",
            "rejectionNote": c.rejection_note if c else None,
            "verified": (c.status == "approved") if c else False,
            "verifiedAt": to_ist_iso(c.verified_at) if c and c.verified_at else None,
        })
        
    has_pending_tasks = any(c.status == "pending_review" for c in enrollment.completions)
    approved_tasks = sum(1 for c in enrollment.completions if c.status == "approved")
    total_tasks = len(enrollment.program.tasks) if enrollment.program else 0
    
    if enrollment.status in ("approved", "rejected"):
        computed_status = enrollment.status
    elif enrollment.status == "master_approved":
        computed_status = "approved"
    elif has_pending_tasks or (total_tasks > 0 and approved_tasks == total_tasks):
        computed_status = "pending"
    else:
        computed_status = "waiting_on_crew"
    
    return {
        "id": enrollment.id,
        "enrollmentId": enrollment.id,
        "vesselName": enrollment.vessel_name,
        "vessel": enrollment.vessel_name,
        "department": enrollment.program.department if enrollment.program else "",
        "submittedAt": to_ist_iso(enrollment.created_at),
        "status": computed_status,
        "decidedAt": to_ist_iso(enrollment.created_at) if computed_status != "pending" else None,
        "candidateName": learner.full_name if learner else "Unknown",
        "candidateRank": learner.rank if learner else "Unknown",
        "candidateCrewId": learner.crew_id if learner else "",
        "programTitle": enrollment.program.title if enrollment and enrollment.program else "Program",
        "learner": {
            "name": learner.full_name if learner else "Unknown",
            "rank": learner.rank if learner else "Unknown",
            "crewId": learner.crew_id if learner else "",
        },
        "program": {
            "title": enrollment.program.title if enrollment and enrollment.program else "Program",
            "subtitle": enrollment.program.subtitle if enrollment and enrollment.program else "",
        },
        "tasks": tasks_out,
        "verifiedCount": approved_tasks,
        "allVerified": total_tasks > 0 and approved_tasks == total_tasks,
    }


@router.get("/api/approver/submissions")
def approver_list_submissions(approver: models.User = Depends(require_vessel_approver), db: Session = Depends(get_db)):
    info = orientation_ranks.vessel_approver_info(approver)
    
    enrollments = (db.query(models.OrientationEnrollment)
                  .join(models.OrientationProgram)
                  .filter(
                      models.OrientationEnrollment.vessel_name == info["vessel"],
                      models.OrientationProgram.department == info["department"],
                  ).all())

    return [orientation_submission_detail(db, e) for e in enrollments]


@router.post("/api/approver/submissions/{submission_id}/tasks/{task_id}/verify")
def approver_verify_task(submission_id: str, task_id: str, req: OrientationVerifyTaskRequest,
                         approver: models.User = Depends(require_vessel_approver), db: Session = Depends(get_db)):
    info = orientation_ranks.vessel_approver_info(approver)
    enrollment = db.get(models.OrientationEnrollment, submission_id)
    if not enrollment or enrollment.vessel_name != info["vessel"] or enrollment.program.department != info["department"]:
        raise HTTPException(404, "Enrollment not found")
        
    completion = (db.query(models.OrientationTaskCompletion)
                  .filter_by(enrollment_id=enrollment.id, task_id=task_id).first())
    if not completion:
        raise HTTPException(404, "Task not found")
        
    if req.action == 'approve':
        completion.status = "approved"
        completion.verified = True
        completion.verified_at = datetime.now(timezone.utc)
        completion.verified_by = approver.id
    elif req.action == 'reject':
        if not req.note:
            raise HTTPException(400, "Rejection note is required")
        completion.status = "rejected"
        completion.rejection_note = req.note
        completion.verified = False
        completion.verified_at = None
        completion.verified_by = None
    else:
        raise HTTPException(400, "Invalid action")
        
    learner = db.get(models.User, enrollment.learner_id)
    task = db.get(models.OrientationTask, task_id)
    if learner and task:
        if req.action == 'approve':
            notify(db, learner.id, "orientation_task_approved", "Task approved",
                   body=f"Your task '{task.title}' was approved by {approver.full_name}.",
                   link="/orientation")
        else:
            notify(db, learner.id, "orientation_task_rejected", "Task rejected",
                   body=f"Your task '{task.title}' was rejected by {approver.full_name}. Note: {req.note}",
                   link="/orientation")
                   
    db.commit()
    return orientation_submission_detail(db, enrollment)


@router.post("/api/approver/submissions/{submission_id}/decide")
def approver_decide_submission(submission_id: str, req: OrientationDecideRequest,
                               approver: models.User = Depends(require_vessel_approver),
                               db: Session = Depends(get_db)):
    info = orientation_ranks.vessel_approver_info(approver)
    enrollment = db.get(models.OrientationEnrollment, submission_id)
    if not enrollment or enrollment.vessel_name != info["vessel"] or enrollment.program.department != info["department"]:
        raise HTTPException(404, "Enrollment not found")

    if req.action == 'approve':
        all_tasks = enrollment.program.tasks
        completions = {c.task_id: c for c in enrollment.completions}
        all_approved = all(completions.get(t.id) and completions[t.id].status == "approved" for t in all_tasks)
        
        if not all_approved:
            raise HTTPException(400, "Cannot approve until all tasks are verified")
            
        enrollment.status = "master_approved"
        db.commit()
        
        if enrollment.learner_id:
            notify(db, enrollment.learner_id, "orientation_pending_admin", "Orientation pending final approval",
                   body=f"All tasks have been approved by {approver.full_name}. Your {enrollment.program.title} is now awaiting final approval from the admin.",
                   link="/orientation")
                   
        admins = db.query(models.User).filter(
            models.User.role.in_(["super_admin", "office_admin", "admin"]),
            models.User.is_active == True
        ).all()
        learner = db.get(models.User, enrollment.learner_id)
        learner_name = learner.full_name if learner else "A crew member"
        for admin in admins:
            notify(db, admin.id, "orientation_master_approved", "Orientation awaiting your approval",
                   body=f"{approver.full_name} (Master/CE) has approved all tasks for {learner_name}'s {enrollment.program.title} on {info['vessel']}. Please give final approval.",
                   link="/admin/orientation-program/enrollments")
                   
        db.commit()
                   
    elif req.action == 'reject':
        enrollment.status = "rejected"
        db.commit()
        
        if enrollment.learner_id:
            notify(db, enrollment.learner_id, "orientation_rejected", "Orientation rejected",
                   body=f"Your {enrollment.program.title} was rejected by {approver.full_name}.",
                   link="/orientation")
    else:
        raise HTTPException(400, "Invalid action")
        
    db.commit()
    return orientation_submission_detail(db, enrollment)


@router.post("/api/admin/orientation/enrollments/{enrollment_id}/decide")
def admin_decide_orientation_enrollment(
    enrollment_id: str,
    req: AdminOrientationDecideRequest,
    admin: models.User = Depends(require_super_admin),
    db: Session = Depends(get_db)
):
    """Super Admin gives the final approval (or rejection) after the Master/CE has approved."""
    enrollment = db.get(models.OrientationEnrollment, enrollment_id)
    if not enrollment:
        raise HTTPException(404, "Enrollment not found")
    if enrollment.status != "master_approved":
        raise HTTPException(400, "Enrollment is not awaiting admin approval")

    if req.action == 'approve':
        enrollment.status = "approved"
        db.commit()
        if enrollment.learner_id:
            notify(db, enrollment.learner_id, "orientation_approved", "Orientation fully approved!",
                   body=f"Congratulations! Your {enrollment.program.title} has been fully approved by the admin. You are now promoted!",
                   link="/orientation")
    elif req.action == 'reject':
        enrollment.status = "rejected"
        db.commit()
        if enrollment.learner_id:
            notify(db, enrollment.learner_id, "orientation_rejected", "Orientation rejected",
                   body=f"Your {enrollment.program.title} was rejected by the admin. Please contact your training officer.",
                   link="/orientation")
    else:
        raise HTTPException(400, "Invalid action")

    db.commit()
    return orientation_submission_detail(db, enrollment)
