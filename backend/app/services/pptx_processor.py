"""PPTX conversion, slide extraction, and background processing service."""
import os
import re
import json
import time
import uuid
import shutil
import subprocess
import threading as _threading
import zipfile as _zipfile
import xml.etree.ElementTree as ET

from app.config import UPLOAD_DIR
from app.database import SessionLocal
from app import models, storage
from app.video import compress_video
from app.notifications_service import notify

_pptx_job_lock = _threading.Lock()


def _pptx_job_path(course_id: str) -> str:
    """Absolute path of the per-course job-state JSON file."""
    return os.path.join(UPLOAD_DIR, course_id, "_pptx_job.json")


def _pptx_job_read(course_id: str) -> dict | None:
    """Return the current job dict, or None if no file exists."""
    try:
        with open(_pptx_job_path(course_id), "r", encoding="utf-8") as fh:
            return json.loads(fh.read())
    except (FileNotFoundError, json.JSONDecodeError):
        return None


def _pptx_job_write(course_id: str, job: dict) -> None:
    """Atomically overwrite the job-state file (tmp + rename)."""
    path = _pptx_job_path(course_id)
    tmp = path + ".tmp"
    os.makedirs(os.path.dirname(path), exist_ok=True)
    with open(tmp, "w", encoding="utf-8") as fh:
        fh.write(json.dumps(job))
    os.replace(tmp, path)


def _safe_unlink(path: str) -> None:
    try:
        os.remove(path)
    except OSError:
        pass


PPTX_SOFFICE_TIMEOUT = int(os.getenv("PPTX_SOFFICE_TIMEOUT", "1800"))

_PPTX_MEDIA_EXT = (
    ".mp4", ".mov", ".avi", ".wmv", ".m4v", ".mkv",
    ".mp3", ".wav", ".m4a", ".wma"
)


def _pptx_render_copy(pptx_path: str, tmp_dir: str) -> str:
    """A copy of the deck with embedded audio/video emptied out, for rendering."""
    try:
        out = os.path.join(tmp_dir, "render_" + os.path.basename(pptx_path))
        stripped = 0
        with _zipfile.ZipFile(pptx_path) as src, _zipfile.ZipFile(out, "w", _zipfile.ZIP_DEFLATED) as dst:
            for item in src.infolist():
                if item.filename.lower().endswith(_PPTX_MEDIA_EXT):
                    dst.writestr(item.filename, b"")
                    stripped += 1
                else:
                    with src.open(item) as f:
                        dst.writestr(item, f.read())
        if not stripped:
            _safe_unlink(out)
            return pptx_path
        before, after = os.path.getsize(pptx_path), os.path.getsize(out)
        print(f"[pptx] render copy: emptied {stripped} media parts "
              f"({before / 1048576:.0f} MB -> {after / 1048576:.0f} MB)")
        return out
    except Exception as e:
        print(f"[pptx] could not build a render copy ({e}); rendering the original")
        return pptx_path


def _pptx_job_set(course_id: str, *, stage: str, pct: int | None = None,
                  message: str = "", error: str | None = None,
                  done: bool = False, added: int | None = None):
    """Update (or create) the on-disk job state for course_id."""
    with _pptx_job_lock:
        job = _pptx_job_read(course_id) or {}
        job.update({
            "stage": stage,
            "message": message,
            "error": error,
            "done": done,
            "updated": time.time(),
        })
        if pct is not None:
            job["pct"] = max(0, min(100, int(pct)))
        if added is not None:
            job["added"] = added
        _pptx_job_write(course_id, job)
    return job


def _next_chapter_order(course) -> int:
    return (max((ch.order for ch in course.chapters), default=-1)) + 1


def _next_chapter_n(course) -> int:
    return (max((ch.n or 0 for ch in course.chapters), default=0)) + 1


def _revoke_course_completion(db, course_id: str):
    """When a new module is added, re-assign the course to users who already completed it,
    and notify all enrolled users."""
    course = db.query(models.Course).filter_by(id=course_id).first()
    if not course:
        return
        
    course_name = course.title
    slug = course.slug if course.slug else course_id
    
    # 1. Reset completion for those who passed
    progress_records = db.query(models.Progress).filter(
        models.Progress.course_id == course_id,
        models.Progress.passed == True
    ).all()
    
    for p in progress_records:
        p.passed = False
        
    # 2. Notify ALL assigned users
    enrollments = db.query(models.Enrollment).filter_by(course_id=course_id).all()
    for e in enrollments:
        notify(
            db, e.learner_id, "info", "Course Updated",
            f"A new module was added to '{course_name}'. Please complete the new content.",
            f"/course/{slug}"
        )


def process_pptx_background(course_id: str, pptx_path: str, original_filename: str, course_dir: str):
    db = SessionLocal()
    _pptx_job_set(course_id, stage="rendering", pct=5,
                  message="Rendering slides (this is the slow part on big decks)…")
    try:
        course = db.query(models.Course).filter(models.Course.id == course_id).first()
        db.close()
        if not course:
            return

        import string
        import random
        # Use a local temp dir to avoid Windows tempfile locking issues with soffice
        tmp = os.path.join(UPLOAD_DIR, "tmp_" + "".join(random.choices(string.ascii_lowercase + string.digits, k=10)))
        os.makedirs(tmp, exist_ok=True)
        
        try:
            soffice_paths = [
                "soffice",
                r"C:\Program Files\LibreOffice\program\soffice.exe",
                r"C:\Program Files (x86)\LibreOffice\program\soffice.exe",
                r"C:\LibreOffice\program\soffice.exe"
            ]
            
            success = False
            last_err = None
            env_tmp = os.path.join(tmp, "soffice_env").replace(os.sep, "/")

            render_src = _pptx_render_copy(pptx_path, tmp)

            for sp in soffice_paths:
                try:
                    subprocess.run(
                        [sp, f"-env:UserInstallation=file:///{env_tmp}", "--headless", "--nologo", "--nofirststartwizard", "--convert-to", "pdf", "--outdir", tmp, render_src],
                        check=True, capture_output=True, timeout=PPTX_SOFFICE_TIMEOUT,
                    )
                    success = True
                    break
                except (subprocess.CalledProcessError, FileNotFoundError, subprocess.TimeoutExpired) as e:
                    last_err = e

            pdf_filename = os.path.splitext(os.path.basename(render_src))[0] + ".pdf"
            pdf_path = os.path.join(tmp, pdf_filename)
            has_pdf = success and os.path.exists(pdf_path)
            
            if not has_pdf:
                print(f"Background PPTX conversion to PDF failed or skipped: {last_err}")
                _pptx_job_set(course_id, stage="rendering", pct=25,
                              message="Slide images unavailable — importing text only.")
            else:
                _pptx_job_set(course_id, stage="extracting", pct=35,
                              message="Extracting slide text and videos…")

            import fitz
            from pptx import Presentation

            try:
                prs = Presentation(pptx_path)
                slide_titles = []
                slide_full_texts = []
                for slide in prs.slides:
                    full_text_parts = []
                    for shape in slide.shapes:
                        if getattr(shape, "has_text_frame", False):
                            full_text_parts.append(shape.text_frame.text.strip())
                        elif getattr(shape, "shape_type", None) == 6:
                            def _get_group_text(shp):
                                res = []
                                for s in shp.shapes:
                                    if getattr(s, "has_text_frame", False):
                                        res.append(s.text_frame.text.strip())
                                    elif getattr(s, "shape_type", None) == 6:
                                        res.extend(_get_group_text(s))
                                return res
                            full_text_parts.extend(_get_group_text(shape))
                    slide_full_texts.append("\n".join(full_text_parts))

                    text = ""
                    title_shape = slide.shapes.title
                    if title_shape and title_shape.has_text_frame:
                        text = title_shape.text_frame.text.strip()
                    
                    if not text:
                        def _get_texts(shapes):
                            res = []
                            for s in shapes:
                                if getattr(s, "shape_type", None) == 6:
                                    res.extend(_get_texts(s.shapes))
                                elif getattr(s, "has_text_frame", False):
                                    t = s.text_frame.text.strip()
                                    if len(t) > 3:
                                        res.append((getattr(s, "top", 0) or 0, t))
                            return res
                        texts = _get_texts(slide.shapes)
                        if texts:
                            texts.sort(key=lambda x: x[0])
                            text = texts[0][1].split("\n")[0].strip()
                            
                    slide_titles.append(text if text else "")
            except Exception as e:
                print(f"Failed to extract text with python-pptx: {e}")
                prs = None
                slide_titles = []
                slide_full_texts = []
            
            import zipfile

            upload_tag = uuid.uuid4().hex[:8]

            slide_part_to_prs_idx: dict[str, int] = {}
            if prs is not None:
                try:
                    for prs_idx, slide in enumerate(prs.slides):
                        part_name = slide.part.partname.lstrip("/")
                        slide_part_to_prs_idx[part_name] = prs_idx
                except Exception as e:
                    print(f"[pptx] could not build slide part map: {e}")

            slide_videos: dict[int, set] = {}
            try:
                with zipfile.ZipFile(pptx_path, "r") as z:
                    namelist_set = set(z.namelist())
                    total_vids = sum(1 for n in z.namelist()
                                     if n.startswith("ppt/media/") and n.lower().endswith(".mp4"))
                    done_vids = 0
                    processed_media: dict[str, str] = {}

                    def _slide_rels_sort_key(n):
                        try:
                            return int(n.split("slide")[2].split(".")[0])
                        except Exception:
                            return 0

                    slide_rels = sorted(
                        [n for n in z.namelist()
                         if n.startswith("ppt/slides/_rels/slide") and n.endswith(".xml.rels")],
                        key=_slide_rels_sort_key
                    )

                    for name in slide_rels:
                        try:
                            slide_num_str = name.split("slide")[2].split(".")[0]

                            slide_part_name = f"ppt/slides/slide{slide_num_str}.xml"
                            if slide_part_name in slide_part_to_prs_idx:
                                slide_idx = slide_part_to_prs_idx[slide_part_name]
                            else:
                                slide_idx = int(slide_num_str) - 1

                            rels_data = z.read(name)
                            root = ET.fromstring(rels_data)
                            ns = {"r": "http://schemas.openxmlformats.org/package/2006/relationships"}
                            for rel in root.findall("r:Relationship", ns):
                                target = rel.get("Target")
                                if target and target.startswith("../media/") and target.lower().endswith(".mp4"):
                                    media_path = "ppt/" + target[3:]
                                    if media_path in namelist_set:
                                        if media_path in processed_media:
                                            slide_videos.setdefault(slide_idx, set()).add(
                                                processed_media[media_path])
                                            continue

                                        basename = os.path.basename(target)
                                        vid_filename = f"{upload_tag}_slide{slide_num_str}_{basename}"
                                        vid_path = os.path.join(course_dir, vid_filename)
                                        with z.open(media_path) as src, open(vid_path, "wb") as vf:
                                            shutil.copyfileobj(src, vf, 1024 * 1024)
                                        _pptx_job_set(
                                            course_id, stage="video",
                                            pct=40 + int(20 * done_vids / max(total_vids, 1)),
                                            message=(f"Compressing video {done_vids + 1} of "
                                                     f"{total_vids} — {basename} "
                                                     f"({os.path.getsize(vid_path) // 1048576} MB). "
                                                     "Large videos take several minutes each."))
                                        final_vid_path = compress_video(vid_path)
                                        done_vids += 1
                                        if final_vid_path == vid_path:
                                            try:
                                                from qtfaststart import processor
                                                processor.process(vid_path, vid_path + ".tmp")
                                                os.replace(vid_path + ".tmp", vid_path)
                                            except Exception:
                                                pass
                                        storage.save(course_id, vid_filename, final_vid_path)
                                        vid_url = f"/api/uploads/{course_id}/{vid_filename}"
                                        processed_media[media_path] = vid_url
                                        slide_videos.setdefault(slide_idx, set()).add(vid_url)
                        except Exception as e:
                            print(f"Error parsing {name}: {e}")
            except Exception as e:
                print(f"Error parsing PPTX zip for videos: {e}")

            doc = fitz.open(pdf_path) if has_pdf else None
            num_slides = len(doc) if doc else (len(prs.slides) if prs else 0)
            
            if num_slides == 0:
                print("No slides found in PPTX or PDF")
                _pptx_job_set(course_id, stage="failed", pct=100, done=True,
                              error="No slides could be read from this file.")
                return

            db = SessionLocal()
            course = db.query(models.Course).filter(models.Course.id == course_id).first()
            if not course:
                _pptx_job_set(course_id, stage="failed", pct=100, done=True,
                              error="Course no longer exists.")
                return

            start_n = max((ch.n or 0 for ch in course.chapters), default=0)
            next_order = _next_chapter_order(course)
            
            for i in range(num_slides):
                n = start_n + i + 1
                
                img_url = None
                if doc:
                    page = doc[i]
                    pix = page.get_pixmap(dpi=150)
                    filename = f"slide{n}.png"
                    img_path = os.path.join(course_dir, filename)
                    pix.save(img_path)
                    storage.save(course_id, filename, img_path)
                    img_url = f"/api/uploads/{course_id}/{filename}"
                
                slide_title = slide_titles[i] if i < len(slide_titles) and slide_titles[i] else f"Slide {n}"
                vids = list(slide_videos.get(i, set()))
                
                full_text = slide_full_texts[i] if i < len(slide_full_texts) else ""
                is_quiz = ("quiz" in slide_title.lower()) or ("A. " in full_text and "B. " in full_text)
                ch_kind = "quiz" if is_quiz else "lesson"
                
                quiz_questions = []
                if is_quiz:
                    lines = full_text.split("\n")
                    prompt_lines = []
                    options = []
                    ans_idx = 0
                    for line in lines:
                        line_s = line.strip()
                        if not line_s:
                            continue
                        
                        ans_match = re.search(r"answer:\s*([A-E])", line_s, re.IGNORECASE)
                        if ans_match:
                            ans_idx = ord(ans_match.group(1).upper()) - ord("A")
                            continue

                        if re.match(r"^[A-E]\.", line_s):
                            options.append(line_s)
                        else:
                            if not options and line_s.lower() != "quiz":
                                prompt_lines.append(line_s)
                    
                    if options:
                        quiz_questions = [models.ChapterQuestion(
                            prompt=" ".join(prompt_lines).strip() or slide_title,
                            options=options,
                            answer=ans_idx,
                            explain="Auto-extracted from slide."
                        )]
                
                ch = models.Chapter(
                    id=f"{course_id}-slide-{n}-{uuid.uuid4().hex[:8]}", course_id=course_id, n=n,
                    title=slide_title,
                    sections=[], videos=vids, order=next_order + i, kind=ch_kind,
                    image=img_url,
                    quiz_questions=quiz_questions
                )
                db.add(ch)
                if num_slides:
                    _pptx_job_set(course_id, stage="slides",
                                  pct=60 + int(35 * (i + 1) / num_slides),
                                  message=f"Building lesson {i + 1} of {num_slides}…")
            
            if doc:
                doc.close()
            if num_slides > 0:
                _revoke_course_completion(db, course_id)
            db.commit()
            _pptx_job_set(course_id, stage="done", pct=100, done=True,
                          added=num_slides,
                          message=f"Imported {num_slides} slides.")
            print(f"Background PPTX processing finished for {course_id}")

        finally:
            try:
                shutil.rmtree(tmp)
            except OSError:
                pass
            db.close()
            
    except Exception as e:
        print(f"Background PPTX processing failed entirely: {e}")
        _pptx_job_set(course_id, stage="failed", pct=100, done=True,
                      error=f"Processing failed: {e}")
        import traceback
        traceback.print_exc()
    finally:
        db.close()
        if os.path.exists(pptx_path):
            try:
                os.remove(pptx_path)
            except OSError:
                pass
